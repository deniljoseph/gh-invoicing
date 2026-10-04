"""
Geometry Home Invoice Management System v8
Multi-company | PostgreSQL (Railway) + SQLite (local) | UAE VAT
"""
import os, io, re, json, time, base64, secrets, zipfile
from datetime import datetime, timedelta, timezone
from functools import wraps
from flask import (Flask, g, render_template, request, redirect, url_for,
                   session, flash, jsonify, send_file, send_from_directory)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from db import get_db, last_insert_id, returning_id, USE_PG

BASE_DIR    = os.path.dirname(os.path.abspath(__file__))
ARCHIVE_DIR = os.path.join(BASE_DIR, "Invoices Archive")
UPLOAD_DIR  = os.path.join(BASE_DIR, "static", "uploads")
IMAGES_DIR  = os.path.join(BASE_DIR, "static", "images")

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', secrets.token_hex(32))
app.config['MAX_CONTENT_LENGTH'] = 512 * 1024 * 1024

for d in [ARCHIVE_DIR, IMAGES_DIR,
          os.path.join(UPLOAD_DIR,"logos"),
          os.path.join(UPLOAD_DIR,"signatures"),
          os.path.join(UPLOAD_DIR,"stamps")]:
    os.makedirs(d, exist_ok=True)

ALLOWED = {'png','jpg','jpeg','gif','webp'}
def allowed_file(fn):
    return '.' in fn and fn.rsplit('.',1)[1].lower() in ALLOWED

# ── DATABASE INIT ──────────────────────────────────────────────────────────
def init_db():
    conn = get_db()

    if USE_PG:
        # PostgreSQL: CREATE TABLE IF NOT EXISTS with pg syntax
        stmts = [
            """CREATE TABLE IF NOT EXISTS users (
                id SERIAL PRIMARY KEY,
                username TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                full_name TEXT, email TEXT,
                role TEXT DEFAULT 'user',
                is_active INTEGER DEFAULT 1,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                last_login TEXT)""",

            """CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY, value TEXT)""",

            """CREATE TABLE IF NOT EXISTS companies (
                id SERIAL PRIMARY KEY,
                code TEXT UNIQUE NOT NULL,
                name TEXT NOT NULL,
                address TEXT, address2 TEXT, address3 TEXT, address4 TEXT,
                telephone TEXT, email TEXT, trn TEXT, website TEXT,
                logo_path TEXT, stamp_path TEXT, letterhead_path TEXT,
                invoice_prefix TEXT, proforma_prefix TEXT,
                is_active INTEGER DEFAULT 1,
                sort_order INTEGER DEFAULT 0,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP)""",

            """CREATE TABLE IF NOT EXISTS clients (
                id SERIAL PRIMARY KEY,
                name TEXT NOT NULL, company_name TEXT,
                trn TEXT, address TEXT, country TEXT,
                telephone TEXT, email TEXT, notes TEXT,
                is_active INTEGER DEFAULT 1,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP)""",

            """CREATE TABLE IF NOT EXISTS signatories (
                id SERIAL PRIMARY KEY,
                name TEXT NOT NULL, position TEXT, image_path TEXT,
                is_active INTEGER DEFAULT 1, is_default INTEGER DEFAULT 0,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP)""",

            """CREATE TABLE IF NOT EXISTS stamps (
                id SERIAL PRIMARY KEY,
                name TEXT NOT NULL, company_id INTEGER,
                image_path TEXT,
                is_active INTEGER DEFAULT 1, is_default INTEGER DEFAULT 0,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP)""",

            """CREATE TABLE IF NOT EXISTS invoices (
                id SERIAL PRIMARY KEY,
                invoice_number TEXT UNIQUE NOT NULL,
                invoice_type TEXT DEFAULT 'TAX',
                company_id INTEGER,
                invoice_date TEXT NOT NULL,
                num_pages TEXT DEFAULT 'One (1)',
                purchase_order TEXT,
                client_id INTEGER,
                client_name TEXT, client_trn TEXT,
                client_address TEXT, client_country TEXT,
                client_telephone TEXT, client_email TEXT,
                bank_name TEXT, bank_iban TEXT,
                bank_account TEXT, bank_swift TEXT,
                currency TEXT DEFAULT 'AED',
                subtotal REAL DEFAULT 0,
                vat_amount REAL DEFAULT 0,
                net_payable REAL DEFAULT 0,
                amount_in_words TEXT,
                payment_terms TEXT, mode_of_payment TEXT, notes TEXT,
                signatory_id INTEGER, signatory_name TEXT, signatory_image TEXT,
                stamp_id INTEGER, stamp_name TEXT, stamp_image TEXT,
                include_signature INTEGER DEFAULT 1,
                include_stamp INTEGER DEFAULT 1,
                pdf_path TEXT, status TEXT DEFAULT 'active',
                created_by INTEGER,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT DEFAULT CURRENT_TIMESTAMP)""",

            """CREATE TABLE IF NOT EXISTS invoice_items (
                id SERIAL PRIMARY KEY,
                invoice_id INTEGER NOT NULL,
                sr_no INTEGER, description TEXT,
                quantity REAL DEFAULT 1, amount REAL DEFAULT 0,
                total_amount REAL DEFAULT 0, tax_rate REAL DEFAULT 5,
                tax_amount REAL DEFAULT 0, total REAL DEFAULT 0)""",

            """CREATE TABLE IF NOT EXISTS audit_logs (
                id SERIAL PRIMARY KEY,
                user_id INTEGER, username TEXT,
                action TEXT, details TEXT, ip_address TEXT,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP)""",
        ]
        for stmt in stmts:
            conn.execute(stmt, ())
        conn.commit()
    else:
        # SQLite
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL,
            full_name TEXT, email TEXT, role TEXT DEFAULT 'user',
            is_active INTEGER DEFAULT 1,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP, last_login TEXT);
        CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT);
        CREATE TABLE IF NOT EXISTS companies (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT UNIQUE NOT NULL, name TEXT NOT NULL,
            address TEXT, address2 TEXT, address3 TEXT, address4 TEXT,
            telephone TEXT, email TEXT, trn TEXT, website TEXT,
            logo_path TEXT, stamp_path TEXT, letterhead_path TEXT,
            invoice_prefix TEXT, proforma_prefix TEXT,
            is_active INTEGER DEFAULT 1, sort_order INTEGER DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS clients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL, company_name TEXT, trn TEXT,
            address TEXT, country TEXT, telephone TEXT, email TEXT, notes TEXT,
            is_active INTEGER DEFAULT 1,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS signatories (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL, position TEXT, image_path TEXT,
            is_active INTEGER DEFAULT 1, is_default INTEGER DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS stamps (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL, company_id INTEGER, image_path TEXT,
            is_active INTEGER DEFAULT 1, is_default INTEGER DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            invoice_number TEXT UNIQUE NOT NULL,
            invoice_type TEXT DEFAULT 'TAX', company_id INTEGER,
            invoice_date TEXT NOT NULL, num_pages TEXT DEFAULT 'One (1)',
            purchase_order TEXT, client_id INTEGER,
            client_name TEXT, client_trn TEXT,
            client_address TEXT, client_country TEXT,
            client_telephone TEXT, client_email TEXT,
            bank_name TEXT, bank_iban TEXT, bank_account TEXT, bank_swift TEXT,
            currency TEXT DEFAULT 'AED',
            subtotal REAL DEFAULT 0, vat_amount REAL DEFAULT 0,
            net_payable REAL DEFAULT 0, amount_in_words TEXT,
            payment_terms TEXT, mode_of_payment TEXT, notes TEXT,
            signatory_id INTEGER, signatory_name TEXT, signatory_image TEXT,
            stamp_id INTEGER, stamp_name TEXT, stamp_image TEXT,
            include_signature INTEGER DEFAULT 1, include_stamp INTEGER DEFAULT 1,
            pdf_path TEXT, status TEXT DEFAULT 'active',
            created_by INTEGER,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS invoice_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            invoice_id INTEGER NOT NULL, sr_no INTEGER,
            description TEXT, quantity REAL DEFAULT 1, amount REAL DEFAULT 0,
            total_amount REAL DEFAULT 0, tax_rate REAL DEFAULT 5,
            tax_amount REAL DEFAULT 0, total REAL DEFAULT 0);
        CREATE TABLE IF NOT EXISTS audit_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER, username TEXT,
            action TEXT, details TEXT, ip_address TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        """)
        # Safe column migrations for existing DBs
        for col in ['address2','address3','address4']:
            try:
                conn.execute(f"ALTER TABLE companies ADD COLUMN {col} TEXT DEFAULT ''", ())
            except Exception:
                pass
        conn.commit()

    # Migrations: proforma <-> tax invoice links, and per-company bank account
    _migs = [("invoices","converted_from_id","INTEGER"),("invoices","converted_from_number","TEXT"),
             ("invoices","converted_to_id","INTEGER"),("invoices","converted_to_number","TEXT"),
             ("invoices","converted_at","TEXT"),
             ("companies","bank_name","TEXT"),("companies","bank_iban","TEXT"),
             ("companies","bank_account","TEXT"),("companies","bank_swift","TEXT"),
             ("users","permissions","TEXT")]
    for _t,_col,_typ in _migs:
        try:
            if USE_PG:
                conn.execute(f"ALTER TABLE {_t} ADD COLUMN IF NOT EXISTS {_col} {_typ}", ())
            else:
                conn.execute(f"ALTER TABLE {_t} ADD COLUMN {_col} {_typ}", ())
            conn.commit()
        except Exception:
            try: conn._conn.rollback()
            except Exception: pass

    # Seed default settings
    defaults = {
        'bank_name':'WIO BUSINESS',
        'bank_iban':'AE430860000009466073611',
        'bank_account':'9466073611',
        'bank_swift':'WIOBAEADXXX',
        'currency':'AED',
        'logo_path':'uploads/logos/GH_LOGO.jpg',
        'footer_path':'images/gh_footer_AE.png',
        'default_payment_terms':'100% Advance Payment',
        'default_mode_of_payment':'Bank Transfer',
        'default_vat_rate':'5',
        'default_company_id':'1',
    }
    for k, v in defaults.items():
        if USE_PG:
            conn.execute("INSERT INTO settings(key,value) VALUES(%s,%s) ON CONFLICT(key) DO NOTHING", (k,v))
        else:
            conn.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", (k,v))

    # Admin user
    if USE_PG:
        row = conn.execute("SELECT id FROM users WHERE username=%s", ('admin',)).fetchone()
    else:
        row = conn.execute("SELECT id FROM users WHERE username=?", ('admin',)).fetchone()
    if not row:
        if USE_PG:
            conn.execute("INSERT INTO users(username,password_hash,full_name,role) VALUES(%s,%s,%s,%s)",
                        ('admin', generate_password_hash('admin123'), 'Administrator', 'admin'))
        else:
            conn.execute("INSERT INTO users(username,password_hash,full_name,role) VALUES(?,?,?,?)",
                        ('admin', generate_password_hash('admin123'), 'Administrator', 'admin'))

    # Seed companies
    companies_data = [
        ('GHM','Geometry Home Furniture Manufacturing L.L.C',
         'L06, L07, L08, WH Phase 3, Dubai Industrial City',
         'PO. Box: 16794,  Phone: +97145578898','','',
         '+97145578898','info@geometry-home.ae','104150505600003','www.geometry-home.ae',
         'uploads/logos/GH_LOGO.jpg','uploads/stamps/GHM_Seal_transparent.png','',
         'GHM-INV','GHM-PINV',1,1),
        ('GHT','Geometry Home for Furniture Trading Co. L.L.C',
         'Shop - GF - 11, Al Barsha Second, Art of Living Mall, Dubai',
         'PO. Box: 16794,  Phone: +97148834020','','',
         '+97148834020','info@geometry-home.ae','104150505600003','www.geometry-home.ae',
         'uploads/logos/GH_LOGO.jpg','uploads/stamps/GHT_Seal_transparent.png','',
         'GHT-INV','GHT-PINV',1,2),
        ('TRA','T R A Geometry Home Technical Services L.L.C S.O.C',
         'A25 \u2013 Hall No. 1, Al Nasr Central, Oud Metha, Dubai',
         'PO. Box: 16794,  Phone: +97145578898','','',
         '+97145578898','info@geometry-home.ae','1533428','www.geometry-home.ae',
         'uploads/logos/GH_LOGO.jpg','uploads/stamps/TRA_Seal_transparent.png','',
         'TRA-INV','TRA-PINV',1,3),
    ]
    for cd in companies_data:
        if USE_PG:
            row = conn.execute("SELECT id FROM companies WHERE code=%s", (cd[0],)).fetchone()
        else:
            row = conn.execute("SELECT id FROM companies WHERE code=?", (cd[0],)).fetchone()
        if not row:
            if USE_PG:
                conn.execute("""INSERT INTO companies(code,name,address,address2,address3,address4,
                    telephone,email,trn,website,logo_path,stamp_path,letterhead_path,
                    invoice_prefix,proforma_prefix,is_active,sort_order)
                    VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""", cd)
            else:
                conn.execute("""INSERT INTO companies(code,name,address,address2,address3,address4,
                    telephone,email,trn,website,logo_path,stamp_path,letterhead_path,
                    invoice_prefix,proforma_prefix,is_active,sort_order)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", cd)

    # Default signatory
    if USE_PG:
        row = conn.execute("SELECT id FROM signatories WHERE name=%s", ('Rashid Khakimov',)).fetchone()
    else:
        row = conn.execute("SELECT id FROM signatories WHERE name=?", ('Rashid Khakimov',)).fetchone()
    if not row:
        if USE_PG:
            conn.execute("INSERT INTO signatories(name,position,image_path,is_active,is_default) VALUES(%s,%s,%s,1,1)",
                        ('Rashid Khakimov','General Manager','uploads/signatures/Rashid_Khakimov_Sign_transparent.png'))
        else:
            conn.execute("INSERT INTO signatories(name,position,image_path,is_active,is_default) VALUES(?,?,?,1,1)",
                        ('Rashid Khakimov','General Manager','uploads/signatures/Rashid_Khakimov_Sign_transparent.png'))

    # Stamps
    stamp_data = [
        ('GHM Official Seal','GHM','uploads/stamps/GHM_Seal_transparent.png',1),
        ('GHT Official Seal','GHT','uploads/stamps/GHT_Seal_transparent.png',0),
        ('TRA Seal',         'TRA','uploads/stamps/TRA_Seal_transparent.png',0),
    ]
    for sname, scode, spath, isdef in stamp_data:
        if USE_PG:
            co_row = conn.execute("SELECT id FROM companies WHERE code=%s", (scode,)).fetchone()
            s_row  = conn.execute("SELECT id FROM stamps WHERE name=%s", (sname,)).fetchone()
        else:
            co_row = conn.execute("SELECT id FROM companies WHERE code=?", (scode,)).fetchone()
            s_row  = conn.execute("SELECT id FROM stamps WHERE name=?", (sname,)).fetchone()
        if co_row and not s_row:
            co_id = co_row['id']
            if USE_PG:
                conn.execute("INSERT INTO stamps(name,company_id,image_path,is_active,is_default) VALUES(%s,%s,%s,1,%s)",
                            (sname, co_id, spath, isdef))
            else:
                conn.execute("INSERT INTO stamps(name,company_id,image_path,is_active,is_default) VALUES(?,?,?,1,?)",
                            (sname, co_id, spath, isdef))

    # One-time: legacy shared bank details (formerly in Settings) now live on the company.
    # Applied only to GHM (the account those details belong to) and only if GHM has no bank account yet.
    try:
        ghm = conn.execute("SELECT id,bank_iban,bank_account FROM companies WHERE code='GHM'").fetchone()
        if ghm and not (ghm['bank_iban'] or '').strip() and not (ghm['bank_account'] or '').strip():
            legacy = {r['key']: r['value'] for r in conn.execute("SELECT key,value FROM settings WHERE key IN ('bank_name','bank_iban','bank_account','bank_swift')").fetchall()}
            if (legacy.get('bank_iban') or legacy.get('bank_account')):
                _p = '%s' if USE_PG else '?'
                conn.execute(f"UPDATE companies SET bank_name={_p},bank_iban={_p},bank_account={_p},bank_swift={_p} WHERE id={_p}",
                             (legacy.get('bank_name',''),legacy.get('bank_iban',''),legacy.get('bank_account',''),legacy.get('bank_swift',''),ghm['id']))
    except Exception:
        app.logger.exception('legacy bank migration skipped')
    conn.commit()
    conn.close()

# ── HELPERS ────────────────────────────────────────────────────────────────
def gset(key, default=''):
    conn = get_db()
    if USE_PG:
        r = conn.execute("SELECT value FROM settings WHERE key=%s", (key,)).fetchone()
    else:
        r = conn.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    conn.close()
    return r['value'] if r else default

def gall():
    conn = get_db()
    rows = conn.execute("SELECT key,value FROM settings", ()).fetchall()
    conn.close()
    return {r['key']: r['value'] for r in rows}

def sset(key, value):
    conn = get_db()
    if USE_PG:
        conn.execute("INSERT INTO settings(key,value) VALUES(%s,%s) ON CONFLICT(key) DO UPDATE SET value=EXCLUDED.value", (key, value))
    else:
        conn.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (key, value))
    conn.commit()
    conn.close()

def get_companies():
    conn = get_db()
    rows = conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order,code", ()).fetchall()
    conn.close()
    return rows

def get_company(cid):
    conn = get_db()
    if USE_PG:
        row = conn.execute("SELECT * FROM companies WHERE id=%s", (cid,)).fetchone()
    else:
        row = conn.execute("SELECT * FROM companies WHERE id=?", (cid,)).fetchone()
    conn.close()
    return dict(row) if row else None

BANK_KEYS = ('bank_name','bank_iban','bank_account','bank_swift')

def bank_for_company(co, s=None):
    """Bank details come only from the company's own record (Companies section)."""
    c = dict(co) if co else {}
    return {k: (c.get(k) or '') for k in BANK_KEYS}

def audit(action, details=''):
    if 'user_id' in session:
        conn = get_db()
        if USE_PG:
            conn.execute("INSERT INTO audit_logs(user_id,username,action,details,ip_address) VALUES(%s,%s,%s,%s,%s)",
                        (session['user_id'], session.get('username'), action, details, request.remote_addr))
        else:
            conn.execute("INSERT INTO audit_logs(user_id,username,action,details,ip_address) VALUES(?,?,?,?,?)",
                        (session['user_id'], session.get('username'), action, details, request.remote_addr))
        conn.commit()
        conn.close()

def next_inv_num(company_id, inv_type='TAX'):
    conn = get_db()
    if USE_PG:
        co = conn.execute("SELECT * FROM companies WHERE id=%s", (company_id,)).fetchone()
        row = conn.execute(
            "SELECT invoice_number FROM invoices WHERE company_id=%s AND invoice_type=%s ORDER BY id DESC LIMIT 1",
            (company_id, inv_type)).fetchone()
    else:
        co = conn.execute("SELECT * FROM companies WHERE id=?", (company_id,)).fetchone()
        row = conn.execute(
            "SELECT invoice_number FROM invoices WHERE company_id=? AND invoice_type=? ORDER BY id DESC LIMIT 1",
            (company_id, inv_type)).fetchone()
    conn.close()
    year = now_dubai().year
    prefix = co['invoice_prefix'] if inv_type == 'TAX' else co['proforma_prefix']
    if row:
        try: num = int(row['invoice_number'].split('-')[-1]) + 1
        except: num = 1
    else:
        num = 1
    return f"{prefix}-{year}-{num:04d}"

def n2w(amount):
    if amount is None: return ''
    try: amount = float(amount)
    except: return ''
    if amount == 0: return 'Zero Dirhams Only'
    ones = ['','One','Two','Three','Four','Five','Six','Seven','Eight','Nine','Ten',
            'Eleven','Twelve','Thirteen','Fourteen','Fifteen','Sixteen','Seventeen','Eighteen','Nineteen']
    tens_w = ['','','Twenty','Thirty','Forty','Fifty','Sixty','Seventy','Eighty','Ninety']
    def sp(n):
        n = int(n)
        if n < 20: return ones[n]
        if n < 100: return tens_w[n//10] + (' ' + ones[n%10] if n%10 else '')
        if n < 1000: return ones[n//100] + ' Hundred' + (' and ' + sp(n%100) if n%100 else '')
        if n < 1000000: return sp(n//1000) + ' Thousand' + (' ' + sp(n%1000) if n%1000 else '')
        return sp(n//1000000) + ' Million' + (' ' + sp(n%1000000) if n%1000000 else '')
    aed = int(amount); fils = round((amount - aed) * 100)
    r = sp(aed) if aed else 'Zero'
    if fils: r += ' and ' + sp(fils) + ' Fils'
    return r + ' Dirhams Only'

def fmt_date(d):
    if not d: return ''
    try:
        if '-' in str(d) and len(str(d)) == 10:
            dt = datetime.strptime(str(d), '%Y-%m-%d')
            return dt.strftime('%d-%m-%Y')
    except: pass
    return str(d)

app.jinja_env.globals['fmt_date'] = fmt_date

# ── Dubai time (UTC+4, no daylight saving) ─────────────────────────────
DUBAI_TZ = timezone(timedelta(hours=4))

def now_dubai():
    """Current Dubai time as a naive datetime (the server itself may run in UTC)."""
    return datetime.now(DUBAI_TZ).replace(tzinfo=None)

def dubai_time(value, fmt='%Y-%m-%d %H:%M'):
    """Show a stored database timestamp (UTC, as written by CURRENT_TIMESTAMP) in Dubai time."""
    if not value: return '—'
    txt = str(value).strip()
    m = re.match(r'(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2}:\d{2})(?:\.\d+)?\s*(Z|[+-]\d{2}(?::?\d{2})?)?$', txt)
    if not m: return txt[:19]
    try:
        dt = datetime.strptime(f"{m.group(1)} {m.group(2)}", '%Y-%m-%d %H:%M:%S')
        off = m.group(3)
        if off and off != 'Z':
            sign = -1 if off[0] == '-' else 1
            digits = off[1:].replace(':', '')
            offset = timedelta(hours=int(digits[:2]), minutes=int(digits[2:4] or 0)) * sign
        else:
            offset = timedelta(0)            # no offset stored -> UTC
        return (dt - offset + timedelta(hours=4)).strftime(fmt)
    except Exception:
        return txt[:19]

app.jinja_env.filters['dubai'] = dubai_time

# ── PERMISSIONS ────────────────────────────────────────────────────────────
# Admin: everything.  Sales User (role 'user'): invoices only, with per-user switches below.
INVOICE_PERMS = {
    'inv_create':   'Create invoices (new, duplicate, convert Proforma to Tax)',
    'inv_edit':     'Edit invoices',
    'inv_delete':   'Delete invoices (and undo conversions)',
    'inv_view_all': "See and manage other users' invoices (otherwise only their own)",
}
DEFAULT_PERMS = ['inv_create', 'inv_edit', 'inv_delete']
SALES_ENDPOINTS = {
    'index', 'login', 'logout', 'change_password', 'static',
    'invoices', 'new_invoice', 'edit_invoice', 'view_invoice', 'delete_invoice',
    'convert_to_tax', 'undo_convert', 'duplicate_invoice', 'download_pdf', 'preview_pdf',
    'api_clients', 'api_stamps', 'api_company', 'api_next_inv_num',
}
ENDPOINT_PERM = {
    'new_invoice': 'inv_create', 'duplicate_invoice': 'inv_create', 'convert_to_tax': 'inv_create',
    'edit_invoice': 'inv_edit',
    'delete_invoice': 'inv_delete', 'undo_convert': 'inv_delete',
}

def user_perms(role, raw):
    if role == 'admin': return set(INVOICE_PERMS)
    if raw is None: return set(DEFAULT_PERMS)
    return {p for p in str(raw).split(',') if p in INVOICE_PERMS}

@app.before_request
def enforce_permissions():
    """Secure by default: anything not on the sales allow-list is admin-only."""
    ep = request.endpoint
    if ep in (None, 'static', 'login') or 'user_id' not in session: return
    conn = get_db()
    u = conn.execute("SELECT id,role,is_active,permissions FROM users WHERE id=?", (session['user_id'],)).fetchone()
    conn.close()
    if not u or not u['is_active']:
        session.clear(); flash('Your account is not active.', 'danger'); return redirect(url_for('login'))
    session['role'] = u['role']
    g.perms = user_perms(u['role'], u['permissions'])
    if u['role'] == 'admin': return
    need = ENDPOINT_PERM.get(ep)
    if ep not in SALES_ENDPOINTS or (need and need not in g.perms):
        flash('You do not have permission to do that.', 'danger')
        return redirect(url_for('invoices'))

@app.context_processor
def inject_perms():
    def can(p): return session.get('role') == 'admin' or p in getattr(g, 'perms', ())
    return dict(can=can)

def can_see_all_invoices():
    return session.get('role') == 'admin' or 'inv_view_all' in getattr(g, 'perms', ())

def invoice_allowed(iid):
    """Admins / 'see all' users: any invoice. Others: only invoices they created."""
    if can_see_all_invoices(): return True
    conn = get_db(); r = conn.execute("SELECT created_by FROM invoices WHERE id=?", (iid,)).fetchone(); conn.close()
    return bool(r) and r['created_by'] == session.get('user_id')

def _no_access():
    flash('You do not have access to that invoice.', 'danger')
    return redirect(url_for('invoices'))

def home_url():
    return url_for('dashboard') if session.get('role') == 'admin' else url_for('invoices')

def login_req(f):
    @wraps(f)
    def d(*a, **k):
        if 'user_id' not in session: return redirect(url_for('login'))
        return f(*a, **k)
    return d

def admin_req(f):
    @wraps(f)
    def d(*a, **k):
        if 'user_id' not in session: return redirect(url_for('login'))
        if session.get('role') != 'admin':
            flash('Admin access required.', 'danger')
            return redirect(url_for('dashboard'))
        return f(*a, **k)
    return d



# ── AUTH ───────────────────────────────────────────────────────────────────
@app.route('/')
def index(): return redirect(home_url() if 'user_id' in session else url_for('login'))

@app.route('/login', methods=['GET','POST'])
def login():
    if request.method=='POST':
        u=request.form.get('username','').strip(); p=request.form.get('password','')
        conn=get_db(); user=conn.execute("SELECT * FROM users WHERE username=? AND is_active=1",(u,)).fetchone(); conn.close()
        if user and check_password_hash(user['password_hash'],p):
            session.clear()
            session['user_id']=user['id']; session['username']=user['username']
            session['full_name']=user['full_name'] or user['username']; session['role']=user['role']
            session.permanent='remember' in request.form
            conn2=get_db(); conn2.execute("UPDATE users SET last_login=CURRENT_TIMESTAMP WHERE id=?",(user['id'],)); conn2.commit(); conn2.close()
            audit('LOGIN'); return redirect(home_url())
        flash('Invalid username or password.','danger')
    return render_template('login.html', s=gall())

@app.route('/logout')
def logout(): audit('LOGOUT'); session.clear(); return redirect(url_for('login'))

@app.route('/change-password', methods=['GET','POST'])
@login_req
def change_password():
    if request.method=='POST':
        cur=request.form.get('current_password',''); nw=request.form.get('new_password',''); cn=request.form.get('confirm_password','')
        if nw!=cn: flash('Passwords do not match.','danger')
        elif len(nw)<6: flash('Minimum 6 characters.','danger')
        else:
            conn=get_db(); user=conn.execute("SELECT * FROM users WHERE id=?",(session['user_id'],)).fetchone()
            if user and check_password_hash(user['password_hash'],cur):
                conn.execute("UPDATE users SET password_hash=? WHERE id=?",(generate_password_hash(nw),session['user_id'])); conn.commit()
                flash('Password changed.','success')
            else: flash('Current password incorrect.','danger')
            conn.close()
    return render_template('change_password.html', s=gall())

# ── DASHBOARD ──────────────────────────────────────────────────────────────
@app.route('/dashboard')
@login_req
def dashboard():
    # Revenue = Tax Invoices only. Proforma invoices are quotes/advance requests, so they are reported separately.
    TAX="COALESCE(invoice_type,'TAX')<>'PROFORMA'"; PRO="invoice_type='PROFORMA'"
    conn=get_db()
    def agg(kind, extra="", params=()):
        r=conn.execute(f"SELECT COUNT(*) c, COALESCE(SUM(net_payable),0) s, COALESCE(SUM(vat_amount),0) v FROM invoices WHERE status='active' AND {kind}{extra}", params).fetchone()
        return {'count':r['c'],'amount':float(r['s']),'vat':float(r['v'])}
    total_cl=conn.execute("SELECT COUNT(*) c FROM clients WHERE is_active=1").fetchone()['c']
    mo=now_dubai().strftime('%Y-%m')
    tax_all=agg(TAX); pro_all=agg(PRO)
    tax_month=agg(TAX," AND invoice_date LIKE ?",(f'{mo}%',)); pro_month=agg(PRO," AND invoice_date LIKE ?",(f'{mo}%',))
    recent=conn.execute("SELECT i.*,c.name as co_name,c.code as co_code FROM invoices i LEFT JOIN companies c ON i.company_id=c.id WHERE i.status='active' ORDER BY i.id DESC LIMIT 10").fetchall()
    logs=conn.execute("SELECT * FROM audit_logs ORDER BY id DESC LIMIT 8").fetchall()
    companies=conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order").fetchall()
    co_stats={}
    for co in companies:
        t=agg(TAX," AND company_id=?",(co['id'],)); p=agg(PRO," AND company_id=?",(co['id'],))
        co_stats[co['id']]={'tax_count':t['count'],'revenue':t['amount'],'pro_count':p['count'],'pro_value':p['amount']}
    conn.close()
    return render_template('dashboard.html', s=gall(), companies=companies,
        tax_all=tax_all, pro_all=pro_all, tax_month=tax_month, pro_month=pro_month,
        total_cl=total_cl, recent=recent, logs=logs, co_stats=co_stats)

# ── INVOICES ───────────────────────────────────────────────────────────────
@app.route('/invoices')
@login_req
def invoices():
    q=request.args.get('q',''); co_id=request.args.get('co',''); inv_type=request.args.get('type','')
    conn=get_db()
    base="SELECT i.*,c.name as co_name,c.code as co_code,(SELECT t.id FROM invoices t WHERE t.id=i.converted_to_id AND t.status='active') AS tax_id FROM invoices i LEFT JOIN companies c ON i.company_id=c.id WHERE i.status='active'"
    params=[]
    if q: base+=" AND (i.invoice_number LIKE ? OR i.client_name LIKE ?)"; params+=[f'%{q}%',f'%{q}%']
    if co_id: base+=" AND i.company_id=?"; params.append(co_id)
    if inv_type: base+=" AND i.invoice_type=?"; params.append(inv_type)
    if not can_see_all_invoices(): base+=" AND i.created_by=?"; params.append(session.get('user_id'))
    base+=" ORDER BY i.id DESC"
    rows=conn.execute(base,params).fetchall()
    companies=conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order").fetchall()
    conn.close()
    return render_template('invoices.html', invoices=rows, q=q, co_id=co_id, inv_type=inv_type, companies=companies, s=gall())

@app.route('/invoices/new', methods=['GET','POST'])
@login_req
def new_invoice():
    s=gall(); conn=get_db()
    co_id=request.args.get('co', s.get('default_company_id','1'))
    inv_type=request.args.get('type','TAX')
    try: co_id=int(co_id)
    except: co_id=1
    companies=conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order").fetchall()
    clients=conn.execute("SELECT * FROM clients WHERE is_active=1 ORDER BY name").fetchall()
    sigs=conn.execute("SELECT * FROM signatories WHERE is_active=1 ORDER BY is_default DESC,name").fetchall()
    stamps=conn.execute("SELECT * FROM stamps WHERE is_active=1 AND company_id=? ORDER BY is_default DESC,name",(co_id,)).fetchall()
    if not stamps: stamps=conn.execute("SELECT * FROM stamps WHERE is_active=1 ORDER BY is_default DESC").fetchall()
    def_sig=conn.execute("SELECT * FROM signatories WHERE is_default=1 AND is_active=1 LIMIT 1").fetchone()
    def_stamp=conn.execute("SELECT * FROM stamps WHERE is_default=1 AND is_active=1 AND company_id=? LIMIT 1",(co_id,)).fetchone()
    co=conn.execute("SELECT * FROM companies WHERE id=?",(co_id,)).fetchone()
    conn.close()
    if request.method=='POST': return _save_inv(None,s)
    return render_template('invoice_form.html', mode='new', inv=None, items=[],
        clients=clients, sigs=sigs, stamps=stamps, def_sig=def_sig, def_stamp=def_stamp,
        inv_num=next_inv_num(co_id,inv_type), s=s, companies=companies,
        selected_co_id=co_id, co=co, inv_type=inv_type,
        bank=bank_for_company(co,s), co_banks={c['id']:bank_for_company(c,s) for c in companies})

@app.route('/invoices/<int:iid>/edit', methods=['GET','POST'])
@login_req
def edit_invoice(iid):
    if not invoice_allowed(iid): return _no_access()
    s=gall(); conn=get_db()
    inv=conn.execute("SELECT * FROM invoices WHERE id=?",(iid,)).fetchone()
    if not inv: flash('Not found.','danger'); conn.close(); return redirect(url_for('invoices'))
    co_id=inv['company_id'] or 1
    items=conn.execute("SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY sr_no",(iid,)).fetchall()
    companies=conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order").fetchall()
    clients=conn.execute("SELECT * FROM clients WHERE is_active=1 ORDER BY name").fetchall()
    sigs=conn.execute("SELECT * FROM signatories WHERE is_active=1 ORDER BY is_default DESC,name").fetchall()
    stamps=conn.execute("SELECT * FROM stamps WHERE is_active=1 ORDER BY is_default DESC,name").fetchall()
    co=conn.execute("SELECT * FROM companies WHERE id=?",(co_id,)).fetchone()
    conn.close()
    if request.method=='POST': return _save_inv(iid,s)
    return render_template('invoice_form.html', mode='edit', inv=inv, items=items,
        clients=clients, sigs=sigs, stamps=stamps, def_sig=None, def_stamp=None,
        inv_num=inv['invoice_number'], s=s, companies=companies,
        selected_co_id=co_id, co=co, inv_type=inv['invoice_type'],
        bank=bank_for_company(co,s), co_banks={c['id']:bank_for_company(c,s) for c in companies})

def _int_or_none(v):
    """Form value -> int, or None for '', 'None', 'null' or anything non-numeric."""
    try: return int(str(v).strip())
    except (TypeError, ValueError): return None

def _save_inv(iid,s):
    held=[]
    try:
        return _save_inv_impl(iid,s,held)
    except Exception as e:
        for _c in held:
            try: _c._conn.rollback()
            except Exception: pass
            try: _c.close()
            except Exception: pass
        app.logger.exception('Saving invoice failed')
        flash(f'Could not save the invoice: {e}','danger')
        return redirect(url_for('edit_invoice',iid=iid) if iid else url_for('new_invoice'))

def _save_inv_impl(iid,s,held):
    f=request.form; conn=get_db(); held.append(conn)
    co_id=int(f.get('company_id',1) or 1)
    inv_type=f.get('invoice_type','TAX')
    sid=_int_or_none(f.get('signatory_id')); sn=''; si=''
    if sid:
        sg=conn.execute("SELECT * FROM signatories WHERE id=?",(sid,)).fetchone()
        if sg: sn=sg['name']; si=sg['image_path'] or ''
    stid=_int_or_none(f.get('stamp_id')); stn=''; sti=''
    if stid:
        st=conn.execute("SELECT * FROM stamps WHERE id=?",(stid,)).fetchone()
        if st: stn=st['name']; sti=st['image_path'] or ''
    sub=float(f.get('subtotal',0) or 0)
    vat=float(f.get('vat_amount',0) or 0)
    net=float(f.get('net_payable',0) or 0)
    # Date: store as YYYY-MM-DD internally
    raw_date=f.get('invoice_date','')
    if raw_date:
        try:
            if '-' in raw_date and len(raw_date)==10:
                parts=raw_date.split('-')
                if len(parts[0])==2: raw_date=f"{parts[2]}-{parts[1]}-{parts[0]}"
        except: pass
    data={
        'invoice_number':f.get('invoice_number',''),
        'invoice_type':inv_type,
        'company_id':co_id,
        'invoice_date':raw_date,
        'num_pages':f.get('num_pages','One (1)'),
        'purchase_order':f.get('purchase_order',''),
        'client_id':_int_or_none(f.get('client_id')),
        'client_name':f.get('client_name',''),
        'client_trn':f.get('client_trn',''),
        'client_address':f.get('client_address',''),
        'client_country':f.get('client_country',''),
        'client_telephone':f.get('client_telephone',''),
        'client_email':f.get('client_email',''),
        'bank_name':f.get('bank_name',''),
        'bank_iban':f.get('bank_iban',''),
        'bank_account':f.get('bank_account',''),
        'bank_swift':f.get('bank_swift',''),
        'currency':f.get('currency','AED'),
        'subtotal':sub,'vat_amount':vat,'net_payable':net,
        'amount_in_words':n2w(net),
        'payment_terms':f.get('payment_terms',''),
        'mode_of_payment':f.get('mode_of_payment',''),
        'notes':f.get('notes',''),
        'signatory_id':sid,'signatory_name':sn,'signatory_image':si,
        'stamp_id':stid,'stamp_name':stn,'stamp_image':sti,
        'include_signature':1 if f.get('include_signature') else 0,
        'include_stamp':1 if f.get('include_stamp') else 0,
        'created_by':session.get('user_id'),
    }
    dup=conn.execute("SELECT id FROM invoices WHERE invoice_number=? AND id<>?",(data['invoice_number'],iid or 0)).fetchone()
    if dup:
        conn.close()
        flash(f"Invoice number {data['invoice_number']} already exists. Please use a different number.",'danger')
        return redirect(url_for('edit_invoice',iid=iid) if iid else url_for('new_invoice',co=co_id,type=inv_type))
    if iid:
        _p='%s' if USE_PG else '?'
        sets=', '.join(f"{k}={_p}" for k in data if k!='created_by')
        vals=[v for k,v in data.items() if k!='created_by']+[iid]
        conn.execute(f"UPDATE invoices SET {sets},updated_at=CURRENT_TIMESTAMP WHERE id={_p}",vals)
        conn.execute("DELETE FROM invoice_items WHERE invoice_id=?",(iid,))
        lbl='INVOICE_EDITED'
    else:
        cols=','.join(data.keys())
        ph=','.join(['%s' if USE_PG else '?']*len(data))
        sql=f"INSERT INTO invoices({cols}) VALUES({ph})"
        if USE_PG:
            sql+=' RETURNING id'
            _cur=conn.execute(sql,list(data.values()))
            _row=_cur.fetchone()
            iid=_row['id'] if _row else None
        else:
            conn.execute(sql,list(data.values()))
            iid=conn.execute('SELECT last_insert_rowid()',()).fetchone()[0]
        lbl='INVOICE_CREATED'
    descs=request.form.getlist('description[]'); qtys=request.form.getlist('quantity[]')
    amts=request.form.getlist('amount[]'); rates=request.form.getlist('tax_rate[]')
    for i,desc in enumerate(descs):
        if not desc.strip(): continue
        def _n(lst,default):
            try: return float(lst[i]) if i<len(lst) and str(lst[i]).strip()!='' else default
            except ValueError: return default
        qty=_n(qtys,1); amt=_n(amts,0); rate=_n(rates,5)
        ta=qty*amt; tx=ta*rate/100; tot=ta+tx
        _ii_sql = "INSERT INTO invoice_items(invoice_id,sr_no,description,quantity,amount,total_amount,tax_rate,tax_amount,total) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s)" if USE_PG else "INSERT INTO invoice_items(invoice_id,sr_no,description,quantity,amount,total_amount,tax_rate,tax_amount,total) VALUES(?,?,?,?,?,?,?,?,?)"
        conn.execute(_ii_sql,
                     (iid,i+1,desc,qty,amt,ta,rate,tx,tot))
    conn.commit()
    pdf=None
    try:
        pdf=gen_pdf(iid)
    except Exception as e:
        app.logger.exception('PDF generation failed for invoice %s', iid)
        flash(f'Invoice saved, but the PDF could not be regenerated: {e}','warning')
    if pdf: conn.execute("UPDATE invoices SET pdf_path=? WHERE id=?",(pdf,iid)); conn.commit()
    conn.close(); audit(lbl,data['invoice_number'])
    flash('Invoice saved successfully!','success')
    return redirect(url_for('view_invoice',iid=iid))

@app.route('/invoices/<int:iid>')
@login_req
def view_invoice(iid):
    if not invoice_allowed(iid): return _no_access()
    conn=get_db()
    inv=conn.execute("SELECT i.*,c.name as co_name,c.code as co_code FROM invoices i LEFT JOIN companies c ON i.company_id=c.id WHERE i.id=?",(iid,)).fetchone()
    items=conn.execute("SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY sr_no",(iid,)).fetchall()
    co=conn.execute("SELECT * FROM companies WHERE id=?",(inv['company_id'],)).fetchone() if inv else None
    linked_tax=_linked_tax(conn,inv) if inv else None
    linked_pro=None
    if inv and inv['converted_from_id']:
        linked_pro=conn.execute("SELECT id,invoice_number,status FROM invoices WHERE id=?",(inv['converted_from_id'],)).fetchone()
    conn.close()
    if not inv: flash('Not found.','danger'); return redirect(url_for('invoices'))
    return render_template('invoice_view.html', inv=inv, items=items, s=gall(), co=co, linked_tax=linked_tax, linked_pro=linked_pro)

@app.route('/invoices/<int:iid>/delete', methods=['POST'])
@login_req
def delete_invoice(iid):
    if not invoice_allowed(iid): return _no_access()
    conn=get_db(); inv=conn.execute("SELECT invoice_number FROM invoices WHERE id=?",(iid,)).fetchone()
    conn.execute("UPDATE invoices SET status='deleted' WHERE id=?",(iid,)); conn.commit(); conn.close()
    if inv: audit('INVOICE_DELETED',inv['invoice_number'])
    flash('Invoice deleted.','success'); return redirect(url_for('invoices'))

def _linked_tax(conn, inv):
    """Active Tax invoice previously created from this proforma (or None)."""
    if not inv['converted_to_id']: return None
    return conn.execute("SELECT * FROM invoices WHERE id=? AND status='active'",(inv['converted_to_id'],)).fetchone()

@app.route('/invoices/<int:iid>/convert-to-tax', methods=['POST'])
@login_req
def convert_to_tax(iid):
    if not invoice_allowed(iid): return _no_access()
    """Create a NEW Tax Invoice from a Proforma. The proforma is kept untouched."""
    conn=get_db(); inv=conn.execute("SELECT * FROM invoices WHERE id=?",(iid,)).fetchone()
    if not inv or inv['status']!='active':
        conn.close(); flash('Invoice not found.','danger'); return redirect(url_for('invoices'))
    if (inv['invoice_type'] or 'TAX')!='PROFORMA':
        conn.close(); flash('Only a Proforma Invoice can be converted.','warning'); return redirect(url_for('view_invoice',iid=iid))
    existing=_linked_tax(conn,inv)
    if existing:
        conn.close(); flash(f"Already converted to Tax Invoice {existing['invoice_number']}.",'warning')
        return redirect(url_for('view_invoice',iid=existing['id']))
    items=conn.execute("SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY sr_no",(iid,)).fetchall()
    conn.close()
    new_no=next_inv_num(inv['company_id'] or 1,'TAX')
    now=now_dubai()
    skip={'id','pdf_path','created_at','updated_at'}
    row={k:inv[k] for k in inv.keys() if k not in skip}
    row.update({'invoice_number':new_no,'invoice_type':'TAX','invoice_date':now.strftime('%Y-%m-%d'),
                'status':'active','converted_from_id':iid,'converted_from_number':inv['invoice_number'],
                'converted_to_id':None,'converted_to_number':None,'converted_at':now.strftime('%Y-%m-%d %H:%M:%S'),
                'created_by':session.get('user_id')})
    conn=get_db()
    cols=','.join(row.keys()); ph=','.join(['?']*len(row))
    sql=f"INSERT INTO invoices({cols}) VALUES({ph})"+(" RETURNING id" if USE_PG else "")
    cur=conn.execute(sql,list(row.values()))
    if USE_PG:
        r=cur.fetchone(); nid=r['id'] if r else None
    else:
        nid=conn.execute('SELECT last_insert_rowid()',()).fetchone()[0]
    for it in items:
        conn.execute("INSERT INTO invoice_items(invoice_id,sr_no,description,quantity,amount,total_amount,tax_rate,tax_amount,total) VALUES(?,?,?,?,?,?,?,?,?)",
            (nid,it['sr_no'],it['description'],it['quantity'],it['amount'],it['total_amount'],it['tax_rate'],it['tax_amount'],it['total']))
    conn.execute("UPDATE invoices SET converted_to_id=?,converted_to_number=?,converted_at=? WHERE id=?",
                 (nid,new_no,now.strftime('%Y-%m-%d %H:%M:%S'),iid))
    conn.commit()
    pdf=gen_pdf(nid)
    if pdf: conn.execute("UPDATE invoices SET pdf_path=? WHERE id=?",(pdf,nid)); conn.commit()
    conn.close(); audit('PROFORMA_CONVERTED_TO_TAX',f'{inv["invoice_number"]} -> {new_no}')
    flash(f'Tax Invoice {new_no} created. Proforma {inv["invoice_number"]} has been kept.','success')
    return redirect(url_for('view_invoice',iid=nid))

@app.route('/invoices/<int:iid>/undo-convert', methods=['POST'])
@login_req
def undo_convert(iid):
    if not invoice_allowed(iid): return _no_access()
    """Undo a conversion: remove the Tax Invoice that was generated; the Proforma stays as it was."""
    conn=get_db(); inv=conn.execute("SELECT * FROM invoices WHERE id=?",(iid,)).fetchone()
    if not inv:
        conn.close(); flash('Invoice not found.','danger'); return redirect(url_for('invoices'))
    if (inv['invoice_type'] or 'TAX')=='PROFORMA':
        pro_id=iid; tax_id=inv['converted_to_id']
    else:
        pro_id=inv['converted_from_id']; tax_id=iid
    tax=conn.execute("SELECT * FROM invoices WHERE id=?",(tax_id,)).fetchone() if tax_id else None
    if tax and not can_see_all_invoices() and tax['created_by'] != session.get('user_id'):
        conn.close(); return _no_access()
    if tax and not tax['converted_from_id']:
        conn.close(); flash('This Tax Invoice was not created from a Proforma.','warning'); return redirect(url_for('view_invoice',iid=iid))
    if tax:
        conn.execute("DELETE FROM invoice_items WHERE invoice_id=?",(tax['id'],))
        conn.execute("DELETE FROM invoices WHERE id=?",(tax['id'],))
    pro=conn.execute("SELECT * FROM invoices WHERE id=?",(pro_id,)).fetchone() if pro_id else None
    if pro:
        conn.execute("UPDATE invoices SET converted_to_id=NULL,converted_to_number=NULL,converted_at=NULL WHERE id=?",(pro_id,))
    conn.commit(); conn.close()
    if tax and tax['pdf_path'] and os.path.exists(tax['pdf_path']):
        try: os.remove(tax['pdf_path'])
        except Exception: pass
    audit('TAX_CONVERSION_UNDONE',f'{tax["invoice_number"] if tax else "?"} removed; proforma {pro["invoice_number"] if pro else "?"} kept')
    flash(f'Conversion undone. Tax Invoice {tax["invoice_number"] if tax else ""} removed.','success')
    return redirect(url_for('view_invoice',iid=pro_id) if pro else url_for('invoices'))

@app.route('/invoices/<int:iid>/duplicate')
@login_req
def duplicate_invoice(iid):
    if not invoice_allowed(iid): return _no_access()
    conn=get_db(); inv=conn.execute("SELECT * FROM invoices WHERE id=?",(iid,)).fetchone()
    items=conn.execute("SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY sr_no",(iid,)).fetchall()
    if not inv: conn.close(); flash('Not found.','danger'); return redirect(url_for('invoices'))
    co_id=inv['company_id'] or 1; inv_type=inv['invoice_type'] or 'TAX'; nn=next_inv_num(co_id,inv_type)
    _ph = lambda n: ','.join(['%s']*n) if USE_PG else ','.join(['?']*n)
    _dup_sql = f"""INSERT INTO invoices(invoice_number,invoice_type,company_id,invoice_date,num_pages,
        purchase_order,client_id,client_name,client_trn,client_address,client_country,
        client_telephone,client_email,bank_name,bank_iban,bank_account,bank_swift,currency,
        subtotal,vat_amount,net_payable,amount_in_words,payment_terms,mode_of_payment,notes,
        signatory_id,signatory_name,signatory_image,stamp_id,stamp_name,stamp_image,
        include_signature,include_stamp,created_by)
        VALUES({_ph(34)})""" + (" RETURNING id" if USE_PG else "")
    _dup_cur=conn.execute(_dup_sql,
        (nn,inv_type,co_id,now_dubai().strftime('%Y-%m-%d'),inv['num_pages'],inv['purchase_order'],
         inv['client_id'],inv['client_name'],inv['client_trn'],inv['client_address'],inv['client_country'],
         inv['client_telephone'],inv['client_email'],inv['bank_name'],inv['bank_iban'],inv['bank_account'],
         inv['bank_swift'],inv['currency'],inv['subtotal'],inv['vat_amount'],inv['net_payable'],
         inv['amount_in_words'],inv['payment_terms'],inv['mode_of_payment'],inv['notes'],
         inv['signatory_id'],inv['signatory_name'],inv['signatory_image'],inv['stamp_id'],
         inv['stamp_name'],inv['stamp_image'],inv['include_signature'],inv['include_stamp'],session.get('user_id')))
    if USE_PG:
        _dup_row=_dup_cur.fetchone()
        nid=_dup_row['id'] if _dup_row else None
    else:
        nid=conn.execute('SELECT last_insert_rowid()',()).fetchone()[0]
    for item in items:
        _ii_sql = "INSERT INTO invoice_items(invoice_id,sr_no,description,quantity,amount,total_amount,tax_rate,tax_amount,total) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s)" if USE_PG else "INSERT INTO invoice_items(invoice_id,sr_no,description,quantity,amount,total_amount,tax_rate,tax_amount,total) VALUES(?,?,?,?,?,?,?,?,?)"
        conn.execute(_ii_sql,
                     (nid,item['sr_no'],item['description'],item['quantity'],item['amount'],item['total_amount'],item['tax_rate'],item['tax_amount'],item['total']))
    conn.commit()
    pdf=gen_pdf(nid)
    if pdf: conn.execute("UPDATE invoices SET pdf_path=? WHERE id=?",(pdf,nid)); conn.commit()
    conn.close(); audit('INVOICE_DUPLICATED',f'{inv["invoice_number"]} -> {nn}')
    flash(f'Duplicated as {nn}','success'); return redirect(url_for('edit_invoice',iid=nid))

@app.route('/invoices/<int:iid>/pdf')
@login_req
def download_pdf(iid):
    if not invoice_allowed(iid): return _no_access()
    conn=get_db(); inv=conn.execute("SELECT * FROM invoices WHERE id=?",(iid,)).fetchone(); conn.close()
    if not inv: flash('Not found.','danger'); return redirect(url_for('invoices'))
    p=inv['pdf_path']
    if not p or not os.path.exists(p): p=gen_pdf(iid)
    if p and os.path.exists(p): return send_file(p,as_attachment=True,download_name=f"{inv['invoice_number']}.pdf")
    flash('PDF not available.','danger'); return redirect(url_for('view_invoice',iid=iid))

@app.route('/invoices/<int:iid>/preview')
@login_req
def preview_pdf(iid):
    if not invoice_allowed(iid): return _no_access()
    conn=get_db(); inv=conn.execute("SELECT * FROM invoices WHERE id=?",(iid,)).fetchone(); conn.close()
    if not inv: return "Not found",404
    p=inv['pdf_path']
    if not p or not os.path.exists(p): p=gen_pdf(iid)
    if p and os.path.exists(p): return send_file(p,mimetype='application/pdf')
    return "PDF generation failed",500

# ── PDF GENERATION ─────────────────────────────────────────────────────────
def gen_pdf(iid):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import (Paragraph, Spacer, Table, TableStyle,
                                    Image as RLImg, HRFlowable, KeepTogether)
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
    from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame

    conn = get_db()
    inv   = conn.execute("SELECT * FROM invoices WHERE id=?", (iid,)).fetchone()
    items = conn.execute(
        "SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY sr_no", (iid,)).fetchall()
    conn.close()
    if not inv: return None

    s   = gall()
    co  = get_company(inv["company_id"])
    if not co: return None
    co  = dict(co)

    inv_type   = inv["invoice_type"] or "TAX"
    type_label = "PROFORMA INVOICE" if inv_type == "PROFORMA" else "TAX INVOICE"

    year     = inv["invoice_date"][:4] if inv["invoice_date"] else str(now_dubai().year)
    ydir     = os.path.join(ARCHIVE_DIR, year)
    os.makedirs(ydir, exist_ok=True)
    pdf_path = os.path.join(ydir, f"{inv['invoice_number']}.pdf")

    W, H     = A4
    SM       = 14 * mm
    TM       = 12 * mm
    FOOTER_H = 14 * mm
    FOOTER_GAP = 5 * mm
    CW       = W - 2 * SM
    content_bottom = FOOTER_H + FOOTER_GAP + 4 * mm

    BK   = colors.black
    WH   = colors.white
    DK   = colors.HexColor("#1a1a1a")
    DK2  = colors.HexColor("#2c2c2c")
    LG   = colors.HexColor("#f5f5f5")
    BGRD = colors.HexColor("#cccccc")
    HL   = colors.HexColor("#e8e8e8")
    ACC  = colors.HexColor("#c8a96e")

    def P(txt, sz=8, bold=False, align=TA_LEFT, col=BK):
        return Paragraph(str(txt or ""), ParagraphStyle("s",
            fontName="Helvetica-Bold" if bold else "Helvetica",
            fontSize=sz, alignment=align, textColor=col,
            leading=sz * 1.35, spaceAfter=0, spaceBefore=0))

    dt = inv["invoice_date"] or ""
    try:
        if "-" in dt and len(dt) == 10:
            parts = dt.split("-")
            if len(parts[0]) == 4:
                dt = f"{parts[2]}-{parts[1]}-{parts[0]}"
    except: pass

    words_text = n2w(float(inv["net_payable"] or 0))

    logo_pdf  = os.path.join(BASE_DIR, "static", "uploads", "logos", "GH_LOGO_pdf.jpg")
    logo_main = os.path.join(BASE_DIR, "static", "uploads", "logos", "GH_LOGO.jpg")
    logo_src  = logo_pdf if os.path.exists(logo_pdf) else logo_main
    logo_img  = None
    if os.path.exists(logo_src):
        try: logo_img = RLImg(logo_src, width=33*mm, height=33*mm)
        except: pass

    footer_src = os.path.join(BASE_DIR, "static", "images", "gh_footer_AE.png")

    co_name  = co.get("name",  "")
    co_trn   = co.get("trn",   "")
    co_all_addr = [x for x in [
        co.get("address",""), co.get("address2",""),
        co.get("address3",""), co.get("address4","")
    ] if x and x.strip()]

    def on_page(canv, doc):
        canv.saveState()
        if os.path.exists(footer_src):
            canv.drawImage(footer_src, SM, FOOTER_GAP,
                           width=CW, height=FOOTER_H,
                           preserveAspectRatio=False)
        canv.restoreState()

    frame = Frame(SM, content_bottom, CW, H - TM - content_bottom,
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)

    elems = []

    # ── LETTERHEAD ──────────────────────────────────────────────────────
    LOGO_W = 34*mm; GAP_W = 6*mm; INFO_W = CW - LOGO_W - GAP_W
    addr_rows = [P(co_name, 13, True, TA_LEFT, DK), Spacer(1, 1.5*mm)]
    for adr in co_all_addr:
        addr_rows.append(P(adr, 9, False, TA_LEFT, DK2))
    info_inner = Table([[item] for item in addr_rows], colWidths=[INFO_W])
    info_inner.setStyle(TableStyle([
        ("TOPPADDING",(0,0),(-1,-1),1),("BOTTOMPADDING",(0,0),(-1,-1),1),
        ("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),0),
    ]))
    lh = Table([[logo_img or P("GH",16,True,TA_CENTER,WH), info_inner]],
               colWidths=[LOGO_W+GAP_W, INFO_W])
    lh.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(0,-1),DK),("ALIGN",(0,0),(0,-1),"CENTER"),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("TOPPADDING",(0,0),(0,-1),4),("BOTTOMPADDING",(0,0),(0,-1),4),
        ("LEFTPADDING",(0,0),(0,-1),4),("RIGHTPADDING",(0,0),(0,-1),4),
        ("LEFTPADDING",(1,0),(1,-1),8),("RIGHTPADDING",(1,0),(1,-1),0),
        ("TOPPADDING",(1,0),(1,-1),4),("BOTTOMPADDING",(1,0),(1,-1),4),
    ]))
    elems.append(lh)
    elems.append(Spacer(1,3*mm))
    elems.append(HRFlowable(width=CW, thickness=0.8, color=BGRD, spaceAfter=3*mm))

    # ── TITLE BAR ───────────────────────────────────────────────────────
    TITLE_W=CW*0.42; DET_L=28*mm; DET_C=5*mm; DET_V=CW-TITLE_W-DET_L-DET_C
    det = Table([
        [P("Invoice No",    8,False,TA_LEFT,WH),P(":",8,False,TA_CENTER,WH),P(inv["invoice_number"],         8,True,TA_LEFT,WH)],
        [P("Billing Date",  8,False,TA_LEFT,WH),P(":",8,False,TA_CENTER,WH),P(dt,                            8,False,TA_LEFT,WH)],
        [P("No. of Pages",  8,False,TA_LEFT,WH),P(":",8,False,TA_CENTER,WH),P(inv["num_pages"] or "One (1)", 8,False,TA_LEFT,WH)],
        [P("",6),P("",6),P("(Including this page)",                                                          6,False,TA_LEFT,WH)],
        [P("Purchase Order",8,False,TA_LEFT,WH),P(":",8,False,TA_CENTER,WH),P(inv["purchase_order"] or "",   8,False,TA_LEFT,WH)],
    ], colWidths=[DET_L,DET_C,DET_V])
    det.setStyle(TableStyle([
        ("TEXTCOLOR",(0,0),(-1,-1),WH),("FONTSIZE",(0,0),(-1,-1),8),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("TOPPADDING",(0,0),(-1,-1),2),("BOTTOMPADDING",(0,0),(-1,-1),2),
        ("LEFTPADDING",(0,0),(-1,-1),2),("RIGHTPADDING",(0,0),(-1,-1),4),
    ]))
    ti = Table([[P(type_label,20,True,TA_LEFT,WH)],[P(f"TRN : {co_trn}",8,False,TA_LEFT,ACC)]],
               colWidths=[TITLE_W])
    ti.setStyle(TableStyle([
        ("TOPPADDING",(0,0),(-1,-1),2),("BOTTOMPADDING",(0,0),(-1,-1),2),
        ("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),0),
    ]))
    tb = Table([[ti,det]], colWidths=[TITLE_W, CW-TITLE_W])
    tb.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,-1),DK),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("LEFTPADDING",(0,0),(0,-1),10),("RIGHTPADDING",(0,0),(0,-1),6),
        ("LEFTPADDING",(1,0),(1,-1),6),("RIGHTPADDING",(1,0),(1,-1),6),
        ("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5),
        ("BOX",(0,0),(-1,-1),0.5,BK),("LINEAFTER",(0,0),(0,-1),0.5,colors.HexColor("#444")),
    ]))
    elems.append(tb)
    elems.append(Spacer(1,2*mm))

    # ── CUSTOMER + BANK ─────────────────────────────────────────────────
    half=CW/2; cL=22*mm; cV=half-cL; bL=22*mm; bV=half-bL
    bk    = bank_for_company(co, s)
    iban  = inv["bank_iban"]    or bk.get("bank_iban","")
    acct  = str(inv["bank_account"] or bk.get("bank_account",""))
    swift = inv["bank_swift"]   or bk.get("bank_swift","")
    curr  = inv["currency"]     or "AED"
    bname = inv["bank_name"]    or bk.get("bank_name","")
    def row(ll,lv,rl,rv): return [P(ll,7.5),P(lv,7.5),P(rl,7.5),P(rv,7.5)]
    # Full company name (Paragraph cells wrap automatically if it is ever too long)
    co_name_short = co_name
    info = [
        [P("Customer Details",8,True),P(""),P("Our Bank Details",8,True),P("")],
        row("Name",   ": "+(inv["client_name"]      or ""),"Name",      ": "+co_name_short),
        row("TRN",    ": "+(inv["client_trn"]        or ""),"IBAN No",   ": "+iban),
        row("Address",": "+(inv["client_address"]    or ""),"Account No",": "+acct),
        row("Country",": "+(inv["client_country"]    or ""),"Swift Code",": "+swift),
        row("Tel No", ": "+(inv["client_telephone"]  or ""),"Currency",  ": "+curr),
        row("Email",  ": "+(inv["client_email"]      or ""),"Bank Name", ": "+bname),
    ]
    it2 = Table(info, colWidths=[cL,cV,bL,bV])
    it2.setStyle(TableStyle([
        ("FONTSIZE",(0,0),(-1,-1),7.5),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("TOPPADDING",(0,0),(-1,-1),1.5),("BOTTOMPADDING",(0,0),(-1,-1),1.5),
        ("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),
        ("FONTNAME",(0,0),(3,0),"Helvetica-Bold"),
        ("SPAN",(0,0),(1,0)),("SPAN",(2,0),(3,0)),
        ("LINEBELOW",(0,0),(1,0),0.5,BGRD),("LINEBELOW",(2,0),(3,0),0.5,BGRD),
        ("BOX",(0,0),(1,-1),0.5,BGRD),("BOX",(2,0),(3,-1),0.5,BGRD),
        ("LINEBEFORE",(2,0),(2,-1),1,BGRD),
    ]))
    elems.append(it2)
    elems.append(Spacer(1,2*mm))

    # ── ITEMS TABLE ─────────────────────────────────────────────────────
    c_sr=10*mm;c_qt=15*mm;c_am=22*mm;c_ta=22*mm;c_tr=12*mm;c_tx=19*mm;c_tt=22*mm
    c_ds=CW-c_sr-c_qt-c_am-c_ta-c_tr-c_tx-c_tt
    cws=[c_sr,c_ds,c_qt,c_am,c_ta,c_tr,c_tx,c_tt]
    def TH(t): return P(t,8,True,TA_CENTER,WH)
    hdr=[TH("Sr."),TH("Description"),TH("Qty"),TH("Amount (AED)"),
         TH("Total Amt (AED)"),TH("Tax %"),TH("Tax Amt"),TH("Total (AED)")]
    irows=[]
    for it in items:
        irows.append([
            P(str(it["sr_no"]),8,False,TA_CENTER),P(it["description"] or "",8),
            P(f"{it['quantity']:,.2f}",8,False,TA_CENTER),
            P(f"{it['amount']:,.2f}",8,False,TA_RIGHT),
            P(f"{it['total_amount']:,.2f}",8,False,TA_RIGHT),
            P(f"{it['tax_rate']:,.1f}%",8,False,TA_CENTER),
            P(f"{it['tax_amount']:,.2f}",8,False,TA_RIGHT),
            P(f"{it['total']:,.2f}",8,False,TA_RIGHT),
        ])
    # No forced empty rows - let content determine table height
    if not irows: irows.append([P("",8) for _ in range(8)])
    itbl=Table([hdr]+irows, colWidths=cws, repeatRows=1, splitByRow=1)
    itbl.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),DK2),("TEXTCOLOR",(0,0),(-1,0),WH),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),8),
        ("ALIGN",(0,0),(-1,0),"CENTER"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("GRID",(0,0),(-1,-1),0.5,BGRD),("ROWBACKGROUNDS",(0,1),(-1,-1),[WH,LG]),
        ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3),
        ("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),
    ]))
    elems.append(itbl)
    elems.append(Spacer(1,1.5*mm))

    # ── TOTALS ──────────────────────────────────────────────────────────
    wl=36*mm; wv=half-wl; tl=half-30*mm; tv=30*mm
    wt=Table([[P("Amount in Words (AED):",8,True)],[P(words_text,8)]], colWidths=[half])
    wt.setStyle(TableStyle([
        ("BOX",(0,0),(-1,-1),0.5,BGRD),("LINEBELOW",(0,0),(0,0),0.5,BGRD),
        ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3),
        ("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
    ]))
    tr=Table([
        [P("SUB-TOTAL AMOUNT (AED)",8,True),  P(f"{inv['subtotal']:,.2f}",8,True,TA_RIGHT)],
        [P("TOTAL VAT AMOUNT (AED)",8,True),  P(f"{inv['vat_amount']:,.2f}",8,True,TA_RIGHT)],
        [P("NET PAYABLE AMOUNT (AED)",9,True),P(f"{inv['net_payable']:,.2f}",9,True,TA_RIGHT)],
    ], colWidths=[tl,tv])
    tr.setStyle(TableStyle([
        ("BOX",(0,0),(-1,-1),0.5,BGRD),("LINEBELOW",(0,0),(-1,0),0.5,BGRD),
        ("LINEBELOW",(0,1),(-1,1),0.5,BGRD),("LINEAFTER",(0,0),(0,-1),0.5,BGRD),
        ("BACKGROUND",(0,2),(-1,2),HL),
        ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4),
        ("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
    ]))
    tot_row=Table([[wt,tr]], colWidths=[half,half])
    tot_row.setStyle(TableStyle([
        ("TOPPADDING",(0,0),(-1,-1),0),("BOTTOMPADDING",(0,0),(-1,-1),0),
        ("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),0),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
    ]))

    # ── PAYMENT TERMS ───────────────────────────────────────────────────
    pr=[[P("Payment Terms & Instructions",8,True)],
        [P(f"Payment Terms  :  {inv['payment_terms'] or ''}",8)],
        [P(f"Mode of Payment  :  {inv['mode_of_payment'] or ''}",8)]]
    if inv["notes"]: pr.append([P(f"Notes  :  {inv['notes']}",8)])
    pt=Table(pr, colWidths=[CW])
    pt.setStyle(TableStyle([
        ("BOX",(0,0),(-1,-1),0.5,BGRD),("LINEBELOW",(0,0),(0,0),0.5,BGRD),
        ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4),
        ("LEFTPADDING",(0,0),(-1,-1),8),("RIGHTPADDING",(0,0),(-1,-1),8),
    ]))

    # ── SIGNATURE + STAMP ───────────────────────────────────────────────
    show_sig  = bool(inv["include_signature"] and inv["signatory_image"])
    show_stmp = bool(inv["include_stamp"]     and inv["stamp_image"])
    # Compact sig/stamp sizes - fits on same page as totals
    sig_w=36*mm; sig_h=16*mm; stmp_w=26*mm; stmp_h=26*mm
    sig_img=Spacer(1,sig_h); stmp_img=Spacer(1,stmp_h)
    if show_sig:
        sp=os.path.join(BASE_DIR,"static",inv["signatory_image"])
        if os.path.exists(sp):
            try: sig_img=RLImg(sp,width=sig_w,height=sig_h)
            except: pass
    if show_stmp:
        sp2=os.path.join(BASE_DIR,"static",inv["stamp_image"])
        if os.path.exists(sp2):
            try: stmp_img=RLImg(sp2,width=stmp_w,height=stmp_h)
            except: pass
    sw=CW*0.55; rw=CW-sw
    st=Table([
        [P(f"For {co_name}",8,True), P("Company Stamp",8,True,TA_CENTER) if show_stmp else P("")],
        [sig_img, stmp_img],
        [P("Authorised Signatory",7.5,True), P("")],
        [P(inv["signatory_name"] or "",7.5), P("")],
    ], colWidths=[sw,rw])
    st.setStyle(TableStyle([
        ("BOX",(0,0),(-1,-1),0.5,BGRD),("LINEAFTER",(0,0),(0,-1),0.5,BGRD),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ALIGN",(1,0),(1,-1),"CENTER"),
        ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3),
        ("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),5),
    ]))

    # KeepTogether ensures totals+terms+sig never split across pages
    elems.append(KeepTogether([
        tot_row, Spacer(1,2*mm),
        pt, Spacer(1,2.5*mm),
        st,
    ]))

    # ── BUILD ───────────────────────────────────────────────────────────
    tmp_path = pdf_path + ".tmp"
    doc=BaseDocTemplate(tmp_path, pagesize=A4,
        leftMargin=SM, rightMargin=SM,
        topMargin=TM, bottomMargin=content_bottom)
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=on_page)])
    doc.build(elems)
    try:
        os.replace(tmp_path, pdf_path)
    except OSError:
        # Old PDF is locked (e.g. still open in a viewer on Windows): keep it and save the new one under a new name
        alt = pdf_path[:-4] + "_" + now_dubai().strftime("%H%M%S") + ".pdf"
        os.replace(tmp_path, alt)
        return alt
    return pdf_path


# ── COMPANIES CRUD ─────────────────────────────────────────────────────────
@app.route('/companies')
@admin_req
def companies():
    conn=get_db(); rows=conn.execute("SELECT * FROM companies ORDER BY sort_order,code").fetchall(); conn.close()
    return render_template('companies.html', companies=rows, s=gall())

@app.route('/companies/new', methods=['GET','POST'])
@admin_req
def new_company():
    if request.method=='POST':
        f=request.form; conn=get_db()
        logo_path='uploads/logos/GH_LOGO.jpg'; stamp_path=''; lh_path=''
        for field,subdir in [('logo','logos'),('stamp','stamps')]:
            if field in request.files:
                fi=request.files[field]
                if fi and allowed_file(fi.filename):
                    fn=secure_filename(fi.filename)
                    fi.save(os.path.join(UPLOAD_DIR,subdir,fn))
                    if field=='logo': logo_path=f'uploads/{subdir}/{fn}'
                    else: stamp_path=f'uploads/{subdir}/{fn}'
        code=f.get('code','').upper().strip()
        _coph = ','.join(['%s']*21) if USE_PG else ','.join(['?']*21)
        _co_sql = f"INSERT INTO companies(code,name,address,address2,address3,address4,telephone,email,trn,website,logo_path,stamp_path,letterhead_path,invoice_prefix,proforma_prefix,is_active,sort_order,bank_name,bank_iban,bank_account,bank_swift) VALUES({_coph})"
        conn.execute(_co_sql,
            (code,f.get('name'),f.get('address'),f.get('address2',''),
             f.get('address3',''),f.get('address4',''),
             f.get('telephone'),f.get('email'),
             f.get('trn'),f.get('website'),logo_path,stamp_path,lh_path,
             f.get('invoice_prefix',f'{code}-INV'),f.get('proforma_prefix',f'{code}-PINV'),
             1,int(f.get('sort_order',99)),
             f.get('bank_name','').strip(),f.get('bank_iban','').strip(),f.get('bank_account','').strip(),f.get('bank_swift','').strip()))
        conn.commit(); conn.close(); audit('COMPANY_ADDED',code); flash('Company added.','success')
        return redirect(url_for('companies'))
    return render_template('company_form.html', co=None, mode='new', s=gall())

@app.route('/companies/<int:cid>/edit', methods=['GET','POST'])
@admin_req
def edit_company(cid):
    conn=get_db(); co=conn.execute("SELECT * FROM companies WHERE id=?",(cid,)).fetchone()
    if request.method=='POST':
        f=request.form
        logo_path=co['logo_path']; stamp_path=co['stamp_path']
        for field,subdir in [('logo','logos'),('stamp','stamps')]:
            if field in request.files:
                fi=request.files[field]
                if fi and fi.filename and allowed_file(fi.filename):
                    fn=secure_filename(fi.filename)
                    fi.save(os.path.join(UPLOAD_DIR,subdir,fn))
                    if field=='logo': logo_path=f'uploads/{subdir}/{fn}'
                    else: stamp_path=f'uploads/{subdir}/{fn}'
        _p2='%s' if USE_PG else '?'
        _cu_sql=f"UPDATE companies SET name={_p2},address={_p2},address2={_p2},address3={_p2},address4={_p2},telephone={_p2},email={_p2},trn={_p2},website={_p2},logo_path={_p2},stamp_path={_p2},invoice_prefix={_p2},proforma_prefix={_p2},is_active={_p2},sort_order={_p2},bank_name={_p2},bank_iban={_p2},bank_account={_p2},bank_swift={_p2} WHERE id={_p2}"
        conn.execute(_cu_sql,
            (f.get('name'),f.get('address'),f.get('address2',''),
             f.get('address3',''),f.get('address4',''),
             f.get('telephone'),f.get('email'),
             f.get('trn'),f.get('website'),logo_path,stamp_path,
             f.get('invoice_prefix'),f.get('proforma_prefix'),
             1 if f.get('is_active') else 0,int(f.get('sort_order',99)),
             f.get('bank_name','').strip(),f.get('bank_iban','').strip(),f.get('bank_account','').strip(),f.get('bank_swift','').strip(),cid))
        conn.commit(); conn.close(); audit('COMPANY_EDITED',f.get('name')); flash('Company updated.','success')
        return redirect(url_for('companies'))
    conn.close()
    return render_template('company_form.html', co=co, mode='edit', s=gall())

@app.route('/companies/<int:cid>/delete', methods=['POST'])
@admin_req
def delete_company(cid):
    conn=get_db()
    inv_count=conn.execute("SELECT COUNT(*) c FROM invoices WHERE company_id=? AND status='active'",(cid,)).fetchone()['c']
    if inv_count>0: flash(f'Cannot delete: {inv_count} active invoices linked.','danger'); conn.close(); return redirect(url_for('companies'))
    conn.execute("DELETE FROM companies WHERE id=?",(cid,)); conn.commit(); conn.close()
    audit('COMPANY_DELETED',f'ID {cid}'); flash('Company deleted.','success'); return redirect(url_for('companies'))

@app.route('/companies/bulk-delete', methods=['POST'])
@admin_req
def bulk_delete_companies():
    ids=request.form.getlist('ids[]')
    conn=get_db()
    deleted=0
    for cid in ids:
        n=conn.execute("SELECT COUNT(*) c FROM invoices WHERE company_id=? AND status='active'",(cid,)).fetchone()['c']
        if n==0:
            conn.execute("DELETE FROM companies WHERE id=?",(cid,)); deleted+=1
    conn.commit(); conn.close()
    flash(f'Deleted {deleted} companies.','success'); return redirect(url_for('companies'))

# ── API ────────────────────────────────────────────────────────────────────
@app.route('/api/clients/search')
@login_req
def api_clients():
    q=request.args.get('q',''); conn=get_db()
    rows=conn.execute("SELECT * FROM clients WHERE is_active=1 AND (name LIKE ? OR company_name LIKE ?) LIMIT 10",(f'%{q}%',f'%{q}%')).fetchall()
    conn.close(); return jsonify([dict(r) for r in rows])

@app.route('/api/stamps/<int:company_id>')
@login_req
def api_stamps(company_id):
    conn=get_db()
    rows=conn.execute("SELECT * FROM stamps WHERE is_active=1 AND company_id=? ORDER BY is_default DESC,name",(company_id,)).fetchall()
    if not rows: rows=conn.execute("SELECT * FROM stamps WHERE is_active=1 ORDER BY is_default DESC").fetchall()
    conn.close(); return jsonify([dict(r) for r in rows])

@app.route('/api/company/<int:company_id>')
@login_req
def api_company(company_id):
    co=get_company(company_id)
    if co: return jsonify(dict(co))
    return jsonify({})

@app.route('/api/next-invoice-number')
@login_req
def api_next_inv_num():
    co_id=int(request.args.get('co_id',1))
    inv_type=request.args.get('type','TAX')
    return jsonify({'number':next_inv_num(co_id,inv_type)})

# ── CLIENTS ────────────────────────────────────────────────────────────────
@app.route('/clients')
@login_req
def clients():
    q=request.args.get('q',''); conn=get_db()
    if q: rows=conn.execute("SELECT * FROM clients WHERE is_active=1 AND (name LIKE ? OR company_name LIKE ? OR trn LIKE ?) ORDER BY name",(f'%{q}%',f'%{q}%',f'%{q}%')).fetchall()
    else: rows=conn.execute("SELECT * FROM clients WHERE is_active=1 ORDER BY name").fetchall()
    conn.close(); return render_template('clients.html', clients=rows, q=q, s=gall())

@app.route('/clients/new', methods=['GET','POST'])
@login_req
def new_client():
    if request.method=='POST':
        f=request.form; conn=get_db()
        conn.execute("INSERT INTO clients(name,company_name,trn,address,country,telephone,email,notes) VALUES(?,?,?,?,?,?,?,?)",
                     (f.get('name'),f.get('company_name'),f.get('trn'),f.get('address'),f.get('country'),f.get('telephone'),f.get('email'),f.get('notes')))
        conn.commit(); conn.close(); audit('CLIENT_ADDED',f.get('name')); flash('Client added.','success')
        return redirect(url_for('clients'))
    return render_template('client_form.html', cl=None, mode='new', s=gall())

@app.route('/clients/<int:cid>/edit', methods=['GET','POST'])
@login_req
def edit_client(cid):
    conn=get_db(); cl=conn.execute("SELECT * FROM clients WHERE id=?",(cid,)).fetchone()
    if request.method=='POST':
        f=request.form
        conn.execute("UPDATE clients SET name=?,company_name=?,trn=?,address=?,country=?,telephone=?,email=?,notes=? WHERE id=?",
                     (f.get('name'),f.get('company_name'),f.get('trn'),f.get('address'),f.get('country'),f.get('telephone'),f.get('email'),f.get('notes'),cid))
        conn.commit(); conn.close(); audit('CLIENT_EDITED',f.get('name')); flash('Updated.','success')
        return redirect(url_for('clients'))
    conn.close(); return render_template('client_form.html', cl=cl, mode='edit', s=gall())

@app.route('/clients/<int:cid>/delete', methods=['POST'])
@login_req
def delete_client(cid):
    conn=get_db(); conn.execute("UPDATE clients SET is_active=0 WHERE id=?",(cid,)); conn.commit(); conn.close()
    audit('CLIENT_DELETED',f'ID {cid}'); flash('Client archived.','success'); return redirect(url_for('clients'))

# ── SIGNATORIES & STAMPS ───────────────────────────────────────────────────
@app.route('/signatories')
@admin_req
def signatories():
    conn=get_db()
    sigs=conn.execute("SELECT * FROM signatories ORDER BY is_default DESC,name").fetchall()
    stamps=conn.execute("SELECT s.*,c.name as co_name,c.code as co_code FROM stamps s LEFT JOIN companies c ON s.company_id=c.id ORDER BY c.sort_order,s.is_default DESC,s.name").fetchall()
    companies=conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order").fetchall()
    conn.close(); return render_template('signatories.html', sigs=sigs, stamps=stamps, companies=companies, s=gall())

@app.route('/signatories/add', methods=['POST'])
@admin_req
def add_signatory():
    name=request.form.get('name',''); pos=request.form.get('position','')
    isdef=1 if request.form.get('is_default') else 0; imgp=''
    if 'image' in request.files:
        fi=request.files['image']
        if fi and allowed_file(fi.filename):
            fn=secure_filename(fi.filename); fi.save(os.path.join(UPLOAD_DIR,'signatures',fn))
            imgp=f'uploads/signatures/{fn}'
    conn=get_db()
    if isdef: conn.execute("UPDATE signatories SET is_default=0")
    conn.execute("INSERT INTO signatories(name,position,image_path,is_default) VALUES(%s,%s,%s,%s)" if USE_PG else "INSERT INTO signatories(name,position,image_path,is_default) VALUES(?,?,?,?)",(name,pos,imgp,isdef))
    conn.commit(); conn.close(); audit('SIGNATORY_ADDED',name); flash('Signatory added.','success')
    return redirect(url_for('signatories'))

@app.route('/signatories/<int:sid>/delete', methods=['POST'])
@admin_req
def delete_signatory(sid):
    conn=get_db(); conn.execute("DELETE FROM signatories WHERE id=?",(sid,)); conn.commit(); conn.close()
    flash('Signatory deleted.','success'); return redirect(url_for('signatories'))

@app.route('/signatories/<int:sid>/toggle', methods=['POST'])
@admin_req
def toggle_signatory(sid):
    conn=get_db(); cur=conn.execute("SELECT is_active FROM signatories WHERE id=?",(sid,)).fetchone()
    if cur: conn.execute("UPDATE signatories SET is_active=? WHERE id=?",(0 if cur['is_active'] else 1,sid))
    conn.commit(); conn.close(); flash('Updated.','success'); return redirect(url_for('signatories'))

@app.route('/signatories/<int:sid>/set-default', methods=['POST'])
@admin_req
def set_default_sig(sid):
    conn=get_db(); conn.execute("UPDATE signatories SET is_default=0"); conn.execute("UPDATE signatories SET is_default=1 WHERE id=?",(sid,)); conn.commit(); conn.close()
    flash('Default set.','success'); return redirect(url_for('signatories'))

@app.route('/stamps/add', methods=['POST'])
@admin_req
def add_stamp():
    name=request.form.get('name',''); co_id=int(request.form.get('company_id',1))
    isdef=1 if request.form.get('is_default') else 0; imgp=''
    if 'image' in request.files:
        fi=request.files['image']
        if fi and allowed_file(fi.filename):
            fn=secure_filename(fi.filename); fi.save(os.path.join(UPLOAD_DIR,'stamps',fn))
            imgp=f'uploads/stamps/{fn}'
    conn=get_db()
    if isdef: conn.execute("UPDATE stamps SET is_default=0 WHERE company_id=?",(co_id,))
    conn.execute("INSERT INTO stamps(name,company_id,image_path,is_default) VALUES(%s,%s,%s,%s)" if USE_PG else "INSERT INTO stamps(name,company_id,image_path,is_default) VALUES(?,?,?,?)",(name,co_id,imgp,isdef))
    conn.commit(); conn.close(); audit('STAMP_ADDED',name); flash('Stamp added.','success')
    return redirect(url_for('signatories'))

@app.route('/stamps/<int:sid>/delete', methods=['POST'])
@admin_req
def delete_stamp(sid):
    conn=get_db(); conn.execute("DELETE FROM stamps WHERE id=?",(sid,)); conn.commit(); conn.close()
    flash('Stamp deleted.','success'); return redirect(url_for('signatories'))

@app.route('/stamps/<int:sid>/toggle', methods=['POST'])
@admin_req
def toggle_stamp(sid):
    conn=get_db(); cur=conn.execute("SELECT is_active FROM stamps WHERE id=?",(sid,)).fetchone()
    if cur: conn.execute("UPDATE stamps SET is_active=? WHERE id=?",(0 if cur['is_active'] else 1,sid))
    conn.commit(); conn.close(); flash('Updated.','success'); return redirect(url_for('signatories'))

@app.route('/stamps/<int:sid>/set-default', methods=['POST'])
@admin_req
def set_default_stamp(sid):
    conn=get_db()
    stmp=conn.execute("SELECT company_id FROM stamps WHERE id=?",(sid,)).fetchone()
    if stmp:
        conn.execute("UPDATE stamps SET is_default=0 WHERE company_id=?",(stmp['company_id'],))
        conn.execute("UPDATE stamps SET is_default=1 WHERE id=?",(sid,))
    conn.commit(); conn.close(); flash('Default stamp set.','success'); return redirect(url_for('signatories'))

# ── USERS ──────────────────────────────────────────────────────────────────
@app.route('/users')
@admin_req
def users():
    conn=get_db(); rows=conn.execute("SELECT * FROM users ORDER BY role,username").fetchall(); conn.close()
    return render_template('users.html', users=rows, s=gall(), perm_defs=INVOICE_PERMS, user_perms=user_perms)

@app.route('/users/new', methods=['GET','POST'])
@admin_req
def new_user():
    if request.method=='POST':
        f=request.form; conn=get_db()
        try:
            _u_sql = "INSERT INTO users(username,password_hash,full_name,email,role,permissions) VALUES(%s,%s,%s,%s,%s,%s)" if USE_PG else "INSERT INTO users(username,password_hash,full_name,email,role,permissions) VALUES(?,?,?,?,?,?)"
            role=f.get('role','user') if f.get('role') in ('admin','user') else 'user'
            conn.execute(_u_sql,
                         (f.get('username'),generate_password_hash(f.get('password','')),f.get('full_name'),f.get('email'),role,
                          ','.join(p for p in INVOICE_PERMS if p in f.getlist('perm'))))
            conn.commit(); audit('USER_CREATED',f.get('username')); flash('User created.','success')
        except Exception as e: flash(f'Error: {e}','danger')
        finally: conn.close()
        return redirect(url_for('users'))
    return render_template('user_form.html', u=None, mode='new', s=gall(), perm_defs=INVOICE_PERMS, u_perms=set(DEFAULT_PERMS))

@app.route('/users/<int:uid>/edit', methods=['GET','POST'])
@admin_req
def edit_user(uid):
    conn=get_db(); u=conn.execute("SELECT * FROM users WHERE id=?",(uid,)).fetchone()
    if request.method=='POST':
        f=request.form
        role=f.get('role') if f.get('role') in ('admin','user') else 'user'
        active=1 if f.get('is_active') else 0
        if uid==session.get('user_id'): role='admin'; active=1      # never lock yourself out
        conn.execute("UPDATE users SET full_name=?,email=?,role=?,is_active=?,permissions=? WHERE id=?",
                     (f.get('full_name'),f.get('email'),role,active,
                      ','.join(p for p in INVOICE_PERMS if p in f.getlist('perm')),uid))
        if f.get('password'): conn.execute("UPDATE users SET password_hash=? WHERE id=?",(generate_password_hash(f.get('password')),uid))
        conn.commit(); conn.close(); audit('USER_EDITED',f.get('full_name')); flash('Updated.','success')
        return redirect(url_for('users'))
    conn.close(); return render_template('user_form.html', u=u, mode='edit', s=gall(), perm_defs=INVOICE_PERMS, u_perms=user_perms(u['role'],u['permissions']))

@app.route('/users/<int:uid>/delete', methods=['POST'])
@admin_req
def delete_user(uid):
    if uid==session.get('user_id'): flash('Cannot delete yourself.','danger'); return redirect(url_for('users'))
    conn=get_db(); conn.execute("DELETE FROM users WHERE id=?",(uid,)); conn.commit(); conn.close()
    audit('USER_DELETED',f'ID {uid}'); flash('User deleted.','success'); return redirect(url_for('users'))

# ── SETTINGS ───────────────────────────────────────────────────────────────
@app.route('/settings', methods=['GET','POST'])
@admin_req
def company_settings():
    if request.method=='POST':
        f=request.form; conn=get_db()
        for k in ['currency',
                  'default_payment_terms','default_mode_of_payment','default_vat_rate','default_company_id']:
            if k in f:
                if USE_PG:
                    conn.execute("INSERT INTO settings(key,value) VALUES(%s,%s) ON CONFLICT(key) DO UPDATE SET value=EXCLUDED.value",(k,f.get(k)))
                else:
                    conn.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)",(k,f.get(k)))
        if 'logo' in request.files:
            fi=request.files['logo']
            if fi and fi.filename and allowed_file(fi.filename):
                fn=secure_filename(fi.filename); fi.save(os.path.join(UPLOAD_DIR,'logos',fn))
                if USE_PG:
                    conn.execute("INSERT INTO settings(key,value) VALUES(%s,%s) ON CONFLICT(key) DO UPDATE SET value=EXCLUDED.value",('logo_path',f'uploads/logos/{fn}'))
                else:
                    conn.execute("INSERT OR REPLACE INTO settings(key,value) VALUES('logo_path',?)",(f'uploads/logos/{fn}',))
        conn.commit(); conn.close(); audit('SETTINGS_UPDATED'); flash('Settings saved.','success')
        return redirect(url_for('company_settings'))
    companies=get_companies()
    return render_template('settings.html', s=gall(), companies=companies)

# ── REPORTS ────────────────────────────────────────────────────────────────
@app.route('/reports')
@login_req
def reports():
    period=request.args.get('period','monthly'); year=request.args.get('year',str(now_dubai().year))
    month=request.args.get('month',str(now_dubai().month).zfill(2)); co_filter=request.args.get('co','')
    inv_type=request.args.get('type','')
    conn=get_db()
    base="SELECT i.*,c.name as co_name,c.code as co_code FROM invoices i LEFT JOIN companies c ON i.company_id=c.id WHERE i.status='active'"
    params=[]
    if co_filter: base+=" AND i.company_id=?"; params.append(co_filter)
    if inv_type: base+=" AND i.invoice_type=?"; params.append(inv_type)
    if period=='daily':
        today=now_dubai().strftime('%Y-%m-%d'); base+=" AND i.invoice_date=?"; params.append(today)
    elif period=='weekly':
        wa=(now_dubai()-timedelta(days=7)).strftime('%Y-%m-%d'); base+=" AND i.invoice_date>=?"; params.append(wa)
    elif period=='monthly': base+=" AND i.invoice_date LIKE ?"; params.append(f'{year}-{month}%')
    elif period=='quarterly':
        q=(int(month)-1)//3; months=[f'{year}-{str(q*3+i+1).zfill(2)}%' for i in range(3)]
        base+=" AND (i.invoice_date LIKE ? OR i.invoice_date LIKE ? OR i.invoice_date LIKE ?)"; params+=months
    else: base+=" AND i.invoice_date LIKE ?"; params.append(f'{year}%')
    rows=conn.execute(base+" ORDER BY i.invoice_date DESC",params).fetchall()
    tax_rows=[r for r in rows if (r['invoice_type'] or 'TAX')!='PROFORMA']
    pro_rows=[r for r in rows if r['invoice_type']=='PROFORMA']
    # Revenue / VAT collected count Tax Invoices only; Proforma is shown separately
    total_sales=sum(r['net_payable'] or 0 for r in tax_rows); total_vat=sum(r['vat_amount'] or 0 for r in tax_rows)
    pro_total=sum(r['net_payable'] or 0 for r in pro_rows)
    chart_data=[]; chart_pro=[]
    for m in range(1,13):
        ms=str(m).zfill(2)
        for kind,out in (("COALESCE(invoice_type,'TAX')<>'PROFORMA'",chart_data),("invoice_type='PROFORMA'",chart_pro)):
            qry=f"SELECT COALESCE(SUM(net_payable),0) s FROM invoices WHERE status='active' AND {kind} AND invoice_date LIKE ?"
            qp=[f'{year}-{ms}%']
            if co_filter: qry+=" AND company_id=?"; qp.append(co_filter)
            out.append(float(conn.execute(qry,qp).fetchone()['s']))
    if inv_type=='PROFORMA': chart_data=[0]*12
    if inv_type=='TAX': chart_pro=[0]*12
    companies=conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order").fetchall()
    conn.close()
    return render_template('reports.html', invoices=rows, total_sales=total_sales, total_vat=total_vat,
        tax_count=len(tax_rows), pro_count=len(pro_rows), pro_total=pro_total,
        period=period, year=year, month=month, chart_data=json.dumps(chart_data), chart_pro=json.dumps(chart_pro),
        s=gall(), companies=companies, co_filter=co_filter, inv_type=inv_type)

@app.route('/audit-log')
@admin_req
def audit_log():
    conn=get_db(); logs=conn.execute("SELECT * FROM audit_logs ORDER BY id DESC LIMIT 500").fetchall(); conn.close()
    return render_template('audit_log.html', logs=logs, s=gall())

BACKUP_TABLES = ['users','settings','companies','clients','signatories','stamps','invoices','invoice_items','audit_logs']

def _table_cols(conn, t):
    """[(column, type)] for a table, on SQLite and PostgreSQL."""
    if USE_PG:
        rows = conn.execute("SELECT column_name AS name, data_type AS type FROM information_schema.columns "
                            "WHERE table_schema=current_schema() AND table_name=? ORDER BY ordinal_position", (t,)).fetchall()
        return [(r['name'], (r['type'] or '').lower()) for r in rows]
    return [(r['name'], (r['type'] or '').lower()) for r in conn.execute(f"PRAGMA table_info({t})").fetchall()]

_ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

@app.route('/backup')
@admin_req
def backup():
    """One-click full backup: every table + uploaded logos/stamps/signatures in a single Excel file."""
    from openpyxl import Workbook
    ts = now_dubai().strftime('%Y%m%d_%H%M%S')
    wb = Workbook(); meta = wb.active; meta.title = '_meta'
    conn = get_db(); counts = {}
    for t in BACKUP_TABLES:
        cols = [c for c,_ in _table_cols(conn, t)]
        rows = conn.execute(f"SELECT * FROM {t}").fetchall()
        ws = wb.create_sheet(t); ws.append(cols)
        for r in rows:
            ws.append([None]*len(cols))
            for ci, c in enumerate(cols, 1):
                v = r[c]
                if isinstance(v, str): v = _ILLEGAL.sub('', v)
                cell = ws.cell(row=ws.max_row, column=ci, value=v)
                if isinstance(v, str): cell.data_type = 's'   # never treat text such as "=abc" as a formula
        counts[t] = len(rows)
    conn.close()
    fs = wb.create_sheet('_files'); fs.append(['path','chunk','data'])
    nfiles = 0
    for root, dirs, files in os.walk(UPLOAD_DIR):
        for fn in sorted(files):
            fp = os.path.join(root, fn); rel = os.path.relpath(fp, BASE_DIR).replace(os.sep, '/')
            with open(fp, 'rb') as fh: b64 = base64.b64encode(fh.read()).decode()
            for i in range(0, max(len(b64),1), 30000):
                fs.append([rel, i//30000, b64[i:i+30000]])
            nfiles += 1
    meta.append(['key','value'])
    for k, v in [('app','GH Invoicing Backup'),('format_version',1),('created',now_dubai().strftime('%Y-%m-%d %H:%M:%S')),
                 ('database','PostgreSQL' if USE_PG else 'SQLite'),('files',nfiles)] + [(f'rows:{t}',n) for t,n in counts.items()]:
        meta.append([k, v])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    audit('BACKUP_CREATED', f'GH_Backup_{ts}.xlsx')
    return send_file(buf, as_attachment=True, download_name=f'GH_Backup_{ts}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# ── shared restore engine (used by Excel and full-ZIP restore) ─────────────
def _coerce_for_column(v, ty):
    if v is None: return None
    if 'int' in ty and isinstance(v, float) and v.is_integer(): return int(v)
    if ('char' in ty or 'text' in ty) and not isinstance(v, str):
        return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)
    if ty in ('real', 'double precision', 'numeric', 'float') and isinstance(v, int) and not isinstance(v, bool): return float(v)
    return v

def _read_sequences(conn):
    """Last id handed out per table, so ids continue exactly as before after a restore."""
    seqs = {}
    try:
        if USE_PG:
            for t in BACKUP_TABLES:
                if 'id' not in dict(_table_cols(conn, t)): continue
                sn = conn.execute("SELECT pg_get_serial_sequence(?, 'id') AS sn", (t,)).fetchone()['sn']
                if not sn: continue
                r = conn.execute(f"SELECT last_value, is_called FROM {sn}").fetchone()
                seqs[t] = int(r['last_value']) if r['is_called'] else int(r['last_value']) - 1
        else:
            for r in conn.execute("SELECT name, seq FROM sqlite_sequence").fetchall():
                if r['name'] in BACKUP_TABLES: seqs[r['name']] = int(r['seq'])
    except Exception:
        app.logger.exception('could not read id sequences')
    return seqs

def _replace_all_tables(conn, tables, sequences=None):
    """Replace the content of every backup table. tables = {name: (columns, rows)}. Caller commits/rolls back."""
    info = {t: _table_cols(conn, t) for t in BACKUP_TABLES}
    for t in reversed(BACKUP_TABLES): conn.execute(f"DELETE FROM {t}")
    counts = {}
    for t in BACKUP_TABLES:
        header, rows = tables[t]
        types = dict(info[t])
        use = [(i, h) for i, h in enumerate(header) if h in types]
        n = 0
        for r in rows:
            if all(v is None for v in r): continue
            vals = [_coerce_for_column(r[i] if i < len(r) else None, types[h]) for i, h in use]
            conn.execute(f"INSERT INTO {t}({','.join(h for _,h in use)}) VALUES({','.join(['?']*len(use))})", vals)
            n += 1
        counts[t] = n
        if 'id' in types:
            maxid = conn.execute(f"SELECT COALESCE(MAX(id),0) AS m FROM {t}").fetchone()['m']
            last = max(int((sequences or {}).get(t, 0)), int(maxid))
            if USE_PG:
                conn.execute(f"SELECT setval(pg_get_serial_sequence('{t}','id'), ?, ?)", (max(last, 1), last >= 1))
            elif sequences and t in sequences:
                try:
                    if conn.execute("SELECT 1 FROM sqlite_sequence WHERE name=?", (t,)).fetchone():
                        conn.execute("UPDATE sqlite_sequence SET seq=? WHERE name=?", (last, t))
                    else:
                        conn.execute("INSERT INTO sqlite_sequence(name,seq) VALUES(?,?)", (t, last))
                except Exception:
                    app.logger.exception('could not restore sqlite sequence for %s', t)
    return counts

def _ensure_upload_dirs():
    for d in (UPLOAD_DIR, ARCHIVE_DIR):
        os.makedirs(d, exist_ok=True)
    for sub in ('logos', 'stamps', 'signatures'):
        os.makedirs(os.path.join(UPLOAD_DIR, sub), exist_ok=True)

# ── FULL ZIP BACKUP ────────────────────────────────────────────────────────
FULL_BACKUP_APP = 'GH Invoicing Full Backup'
_FILE_ROOTS = (('uploads', lambda: UPLOAD_DIR), ('archive', lambda: ARCHIVE_DIR))

def _json_default(o):
    from decimal import Decimal
    if isinstance(o, Decimal): return float(o)
    if hasattr(o, 'isoformat'): return o.isoformat(sep=' ')
    return str(o)

def _sha256_bytes(b):
    import hashlib; return hashlib.sha256(b).hexdigest()

def _pdf_path_out(v):
    """Store archive PDF paths relative to the archive folder so they work on any server."""
    if isinstance(v, str) and v:
        root = os.path.normpath(ARCHIVE_DIR) + os.sep
        if os.path.normpath(v).startswith(root):
            return '@archive/' + os.path.relpath(os.path.normpath(v), ARCHIVE_DIR).replace(os.sep, '/')
    return v

def _pdf_path_in(v):
    if isinstance(v, str) and v.startswith('@archive/'):
        return os.path.join(ARCHIVE_DIR, *v[len('@archive/'):].split('/'))
    return v

def _build_full_backup(fobj):
    """Write the complete backup (all tables exactly + every uploaded/archived file + checksums) into fobj."""
    import zipfile
    manifest = {'app': FULL_BACKUP_APP, 'format_version': 1,
                'created': now_dubai().strftime('%Y-%m-%d %H:%M:%S'),
                'database': 'PostgreSQL' if USE_PG else 'SQLite', 'tables': {}, 'sequences': {}, 'files': []}
    with zipfile.ZipFile(fobj, 'w', zipfile.ZIP_DEFLATED) as z:
        conn = get_db()
        for t in BACKUP_TABLES:
            cols = [c for c, _ in _table_cols(conn, t)]
            rows = conn.execute(f"SELECT * FROM {t}").fetchall()
            data = []
            for r in rows:
                data.append([_pdf_path_out(r[c]) if (t == 'invoices' and c == 'pdf_path') else r[c] for c in cols])
            payload = json.dumps({'table': t, 'columns': cols, 'rows': data}, ensure_ascii=False, default=_json_default).encode('utf-8')
            z.writestr(f'database/{t}.json', payload)
            manifest['tables'][t] = {'rows': len(rows), 'sha256': _sha256_bytes(payload)}
        manifest['sequences'] = _read_sequences(conn)
        conn.close()
        for label, getroot in _FILE_ROOTS:
            root = getroot()
            for dp, dn, fns in os.walk(root):
                dn.sort()
                for fn in sorted(fns):
                    fp = os.path.join(dp, fn)
                    if fn.endswith('.restore_tmp'): continue
                    rel = os.path.relpath(fp, root).replace(os.sep, '/')
                    arc = f'files/{label}/{rel}'
                    with open(fp, 'rb') as fh: data = fh.read()
                    mt = os.path.getmtime(fp)
                    dt = time.localtime(mt)[:6]
                    zi = zipfile.ZipInfo(arc, date_time=dt if dt[0] >= 1980 else (1980, 1, 1, 0, 0, 0))
                    zi.compress_type = zipfile.ZIP_DEFLATED
                    z.writestr(zi, data)
                    manifest['files'].append({'path': arc, 'size': len(data), 'sha256': _sha256_bytes(data), 'mtime': mt})
        z.writestr('manifest.json', json.dumps(manifest, indent=1))
        z.writestr('README.txt',
            "GH Invoicing - full backup\n"
            f"Created: {manifest['created']} (Dubai time)\n\n"
            "database/*.json   every table, all columns, exactly as stored (ids included)\n"
            "files/uploads/    logos, stamps, signatures, letterheads\n"
            "files/archive/    generated invoice PDFs\n"
            "manifest.json     row counts and SHA-256 checksums of everything above\n\n"
            "To restore: sign in as admin > Restore Backup > choose this .zip file.\n"
            "Restoring replaces ALL current data with the contents of this file.\n")
    return manifest

def _save_prestore_snapshot():
    """Safety copy of the current state, kept on the server in ./Backups (newest 3)."""
    try:
        d = os.path.join(BASE_DIR, 'Backups'); os.makedirs(d, exist_ok=True)
        name = f"pre-restore_{now_dubai().strftime('%Y%m%d_%H%M%S')}.zip"
        with open(os.path.join(d, name), 'wb') as fh: _build_full_backup(fh)
        old = sorted(x for x in os.listdir(d) if x.startswith('pre-restore_') and x.endswith('.zip'))
        for x in old[:-3]:
            try: os.remove(os.path.join(d, x))
            except OSError: pass
        return name
    except Exception:
        app.logger.exception('pre-restore snapshot failed')
        return None

@app.route('/backup/zip')
@admin_req
def backup_zip():
    """One-click complete backup (database + all files) as a single ZIP."""
    import tempfile
    tmp = tempfile.TemporaryFile()
    man = _build_full_backup(tmp); tmp.seek(0)
    audit('FULL_BACKUP_CREATED', f"{sum(t['rows'] for t in man['tables'].values())} rows, {len(man['files'])} files")
    return send_file(tmp, as_attachment=True, download_name=f"GH_FullBackup_{now_dubai().strftime('%Y%m%d_%H%M%S')}.zip",
                     mimetype='application/zip')

def _restore_full_zip(fi):
    """Verify first (nothing is touched unless every checksum matches), then restore database + files exactly."""
    import zipfile, hashlib, shutil
    try:
        z = zipfile.ZipFile(fi.stream)
        bad = z.testzip()
        if bad: raise ValueError(f'Corrupt file inside the zip: {bad}')
        man = json.loads(z.read('manifest.json'))
        if man.get('app') != FULL_BACKUP_APP: raise ValueError('This is not a GH Invoicing full backup.')
        tables = {}
        for t in BACKUP_TABLES:
            meta = (man.get('tables') or {}).get(t)
            if meta is None: raise ValueError(f'Table "{t}" is missing from the backup.')
            raw = z.read(f'database/{t}.json')
            if _sha256_bytes(raw) != meta['sha256']: raise ValueError(f'Checksum mismatch for table "{t}" - the backup is damaged.')
            d = json.loads(raw)
            if len(d['rows']) != meta['rows']: raise ValueError(f'Row count mismatch for table "{t}".')
            cols = d['columns']
            if t == 'invoices' and 'pdf_path' in cols:
                pi = cols.index('pdf_path')
                for r in d['rows']: r[pi] = _pdf_path_in(r[pi])
            tables[t] = (cols, d['rows'])
        entries = {}
        for e in man.get('files', []):
            p = e['path']; parts = p.split('/')
            if len(parts) < 3 or parts[0] != 'files' or parts[1] not in ('uploads', 'archive') or '..' in parts or p.startswith('/') or '\\' in p:
                raise ValueError(f'Unsafe path in backup: {p}')
            h = hashlib.sha256(); size = 0
            with z.open(p) as src:
                for chunk in iter(lambda: src.read(1 << 20), b''):
                    h.update(chunk); size += len(chunk)
            if h.hexdigest() != e['sha256'] or size != e['size']: raise ValueError(f'Checksum mismatch for file {p}.')
            entries[p] = e
    except Exception as ex:
        flash(f'Cannot restore: {ex}', 'danger'); return redirect(url_for('restore_backup'))

    snapshot = _save_prestore_snapshot()
    conn = get_db()
    try:
        counts = _replace_all_tables(conn, tables, man.get('sequences'))
        conn.commit()
    except Exception as ex:
        try: conn._conn.rollback()
        except Exception: pass
        conn.close(); app.logger.exception('zip restore failed (database)')
        flash(f'Restore failed, nothing was changed: {ex}', 'danger'); return redirect(url_for('restore_backup'))
    conn.close()

    try:
        for label, getroot in _FILE_ROOTS:
            root = getroot(); tmp = root + '.restore_tmp'; old = root + '.restore_old'
            for p in (tmp, old):
                if os.path.exists(p): shutil.rmtree(p, ignore_errors=True)
            os.makedirs(tmp)
            prefix = f'files/{label}/'
            for p, e in entries.items():
                if not p.startswith(prefix): continue
                dest = os.path.join(tmp, *p[len(prefix):].split('/'))
                os.makedirs(os.path.dirname(dest), exist_ok=True)
                with z.open(p) as src, open(dest, 'wb') as out: shutil.copyfileobj(src, out)
                if e.get('mtime'): os.utime(dest, (e['mtime'], e['mtime']))
            if os.path.exists(root): os.replace(root, old)
            os.replace(tmp, root)
            shutil.rmtree(old, ignore_errors=True)
        _ensure_upload_dirs()
    except Exception as ex:
        app.logger.exception('zip restore failed (files)')
        flash(f'Database restored, but writing files failed: {ex}. A safety copy is in the Backups folder.', 'danger')
        return redirect(url_for('login'))
    app.logger.info('Full backup restored: %s', counts)
    session.clear()
    flash(f"Full backup restored exactly ({counts.get('invoices',0)} invoices, {counts.get('companies',0)} companies, "
          f"{len(entries)} files). Please log in again." + (f" A safety copy of the previous data was saved as Backups/{snapshot}." if snapshot else ''), 'success')
    return redirect(url_for('login'))

@app.route('/restore', methods=['GET','POST'])
@admin_req
def restore_backup():
    if request.method == 'GET':
        return render_template('restore.html', s=gall())
    fi = request.files.get('backup_file')
    name = (fi.filename or '').lower() if fi else ''
    if name.endswith('.zip'): return _restore_full_zip(fi)
    from openpyxl import load_workbook
    if not fi or not name.endswith('.xlsx'):
        flash('Please choose a GH backup file (.zip full backup or .xlsx).', 'danger'); return redirect(url_for('restore_backup'))
    try:
        wb = load_workbook(fi, data_only=True)
        if '_meta' not in wb.sheetnames or any(t not in wb.sheetnames for t in BACKUP_TABLES):
            raise ValueError('This is not a valid GH Invoicing backup file.')
        mk = {r[0]: r[1] for r in wb['_meta'].iter_rows(min_row=2, values_only=True) if r and r[0]}
        if mk.get('app') != 'GH Invoicing Backup': raise ValueError('This is not a valid GH Invoicing backup file.')
    except Exception as e:
        flash(f'Cannot read backup: {e}', 'danger'); return redirect(url_for('restore_backup'))

    conn = get_db()
    try:
        tables = {}
        for t in BACKUP_TABLES:
            rows = wb[t].iter_rows(values_only=True)
            header = list(next(rows, None) or ())
            tables[t] = (header, rows)
        restored = _replace_all_tables(conn, tables)
        conn.commit()
    except Exception as e:
        try: conn._conn.rollback()
        except Exception: pass
        conn.close(); app.logger.exception('restore failed')
        flash(f'Restore failed, nothing was changed: {e}', 'danger'); return redirect(url_for('restore_backup'))
    conn.close()

    # restore uploaded images (logos, stamps, signatures)
    chunks = {}
    for r in wb['_files'].iter_rows(min_row=2, values_only=True):
        if r and r[0]: chunks.setdefault(r[0], []).append((int(r[1] or 0), r[2] or ''))
    nf = 0; up_root = os.path.abspath(UPLOAD_DIR)
    for rel, parts in chunks.items():
        dest = os.path.abspath(os.path.join(BASE_DIR, rel))
        if not dest.startswith(up_root + os.sep): continue
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        with open(dest, 'wb') as fh: fh.write(base64.b64decode(''.join(d for _, d in sorted(parts))))
        nf += 1
    audit('BACKUP_RESTORED', f"{fi.filename}: " + ', '.join(f'{t}={n}' for t, n in restored.items()))
    session.clear()
    flash(f"Backup restored ({restored.get('invoices',0)} invoices, {restored.get('companies',0)} companies, {nf} image files). Please log in again.", 'success')
    return redirect(url_for('login'))

@app.route('/archive')
@login_req
def archive():
    years=[]
    if os.path.exists(ARCHIVE_DIR):
        for y in sorted(os.listdir(ARCHIVE_DIR),reverse=True):
            yd=os.path.join(ARCHIVE_DIR,y)
            if os.path.isdir(yd):
                fls=sorted(os.listdir(yd),reverse=True)
                years.append({'year':y,'files':fls,'count':len(fls)})
    return render_template('archive.html', years=years, archive_dir=ARCHIVE_DIR, s=gall())

@app.route('/archive/<year>/<fn>')
@login_req
def dl_archive(year,fn):
    return send_from_directory(os.path.join(ARCHIVE_DIR,year),fn,as_attachment=True)

# ── Auto-initialize DB on startup (works with gunicorn + Railway) ──────────
# Called at module import time so gunicorn --preload runs this ONCE before
# forking workers. Tables are created before any request is handled.
import logging as _logging
_logging.basicConfig(level=_logging.INFO)
_logger = _logging.getLogger(__name__)
_logger.info("Running init_db() at startup...")
init_db()
_logger.info("init_db() completed successfully.")

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    print("=" * 55)
    print("  Geometry Home Invoice System")
    print(f"  Running at: http://localhost:{port}")
    print("  Login: admin / admin123")
    print("=" * 55)
    app.run(debug=False, host='0.0.0.0', port=port)
